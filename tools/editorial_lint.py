#!/usr/bin/env python3
"""Validate reference completeness, editorial copy, and fixed presentation budgets."""
from __future__ import annotations
from pathlib import Path
import json, re, sys, urllib.parse
from bs4 import BeautifulSoup
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
KEYS=['devops','security','finops','architecture']
PPTX={'devops':'aws-devops-agent-presentation.pptx','security':'aws-security-agent-presentation.pptx','finops':'aws-finops-agent-presentation.pptx','architecture':'aws-frontier-agents-common-architecture-presentation.pptx'}
REPORTS={'devops':'devops/validated-report.html','security':'security/report.html','finops':'finops/report.html','architecture':'architecture/common-patterns.html'}
FORBIDDEN=r'(?i)\b(revolutionary|transformative|game[- ]changing|seamless|effortless|best[- ]in[- ]class|zero[- ]touch|cutting[- ]edge|next[- ]generation|evidence accelerator|unlock potential)\b'
errors=[];summary=[]
annotations=json.loads((ROOT/'content/format-annotations.json').read_text())
for k in KEYS:
    evidence=json.loads((ROOT/f'content/{k}-evidence.json').read_text());capabilities=evidence.get('capabilities',evidence.get('patterns',[]));slides=json.loads((ROOT/f'content/{k}-presentation.json').read_text())
    story=json.loads((ROOT/f'content/narratives/{k}.json').read_text())
    if len(annotations['stage_owners'][k])!=len(story['stages']):errors.append(k+' role count mismatch')
    if set(annotations['capability_records'][k])!={c['id'] for c in capabilities}:errors.append(k+' missing/extra capability annotations')
    for c in capabilities:
        for field in ('setup','check'):
            if not c.get(field):errors.append(k+': '+c['id']+' missing '+field)
        if not c.get('claim',c.get('rule')):errors.append(k+': '+c['id']+' missing documented function/rule')
        if not c.get('limit',c.get('boundary')):errors.append(k+': '+c['id']+' missing limit/boundary')
        if len(annotations['capability_records'][k][c['id']]['record'])<45:errors.append(k+': '+c['id']+' retention record too vague')
    for slide in slides['slides']:
        for f in ('title','takeaway'):
            if re.search(FORBIDDEN,slide[f]):errors.append(k+': promotional '+f+': '+slide[f])
    if re.search(FORBIDDEN,story['strapline']):errors.append(k+' promotional strapline')
    if sum(s['seconds'] for s in slides['slides'])!=slides['minutes']*60:errors.append(k+' core time mismatch')
    if sum(s['seconds'] for s in slides['slides'] if s['kind']=='demo')!=slides['demo']*60:errors.append(k+' demo time mismatch')
    soup=BeautifulSoup((ROOT/REPORTS[k]).read_text(),'html.parser')
    refslides=soup.select('.slide.reference-slide')
    if len(refslides)!=len(capabilities):errors.append(k+' HTML reference count mismatch')
    if len(soup.select('.slide[data-kind="timed"]'))!=len(slides['slides']):errors.append(k+' timed core mismatch')
    if len(soup.select('.cap'))!=len(capabilities):errors.append(k+' HTML capability count mismatch')
    for c in soup.select('.cap'):
        if len(c.select('.cap-section'))<4:errors.append(k+' incomplete capability '+c['id'])
        if not c.select('.evidence-strip a'):errors.append(k+' missing cited capability '+c['id'])
    actual=Presentation(ROOT/'presentations'/PPTX[k])
    if len(actual.slides)!=len(slides['slides'])+len(capabilities)+1:errors.append(k+' PowerPoint count mismatch')
    summary.append({'guide':k,'timed_slides':len(slides['slides']),'reference_slides':len(capabilities),'pptx_slides':len(actual.slides),'minutes':slides['minutes'],'demo_minutes':slides['demo']})
result={'revision':'2026-10-21-three-agent-scope','guides':summary,'review_records':sum(x['reference_slides'] for x in summary),'errors':errors,'boundary':'Editorial structure and linked source identifiers verified. No new external documentation verification or live AWS execution.'}
(ROOT/'qa'/'editorial-lint.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result,indent=2))
if errors:sys.exit(1)
