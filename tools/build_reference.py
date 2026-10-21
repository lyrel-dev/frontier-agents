#!/usr/bin/env python3
"""Render the three product references. No historical copies or ZIP bundles are made."""
from __future__ import annotations
import argparse
import hashlib
import html
import json
import subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
KEYS=('devops','security','finops')
REPORT={'devops':'validated-report.html','security':'report.html','finops':'report.html'}
PPT={'devops':'aws-devops-agent-presentation.pptx','security':'aws-security-agent-presentation.pptx','finops':'aws-finops-agent-presentation.pptx'}

def E(value): return html.escape(str(value),quote=True)
def write(name,content):
    p=ROOT/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(content,encoding='utf-8')
def stamp(seconds):
    return f'{seconds//60}:{seconds%60:02d}'
def load(key): return json.loads((ROOT/'reference'/f'{key}.json').read_text(encoding='utf-8'))
def refs(d,ids,prefix=''):
    return '<div class="evidence">Official documentation: '+ ' · '.join(f'<a href="{prefix}#source-{E(i)}">{E(d["sources"][i]["title"])}</a>' for i in dict.fromkeys(ids))+'</div>' if ids else ''
def p(text,cls=''): return f'<p class="{cls}">{E(text)}</p>'
def block(label,text,cls=''):return f'<div class="field {cls}"><h4>{E(label)}</h4>{p(text)}</div>'
def table(headers,rows):
    return '<div class="table-wrap"><table><thead><tr>'+''.join(f'<th scope="col">{E(x)}</th>' for x in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join(f'<td>{E(x)}</td>' for x in row)+'</tr>' for row in rows)+'</tbody></table></div>'

CSS='''
:root{--ink:#182b40;--muted:#536477;--navy:#11283e;--accent:#006d76;--paper:#fff;--wash:#f3f6f8;--line:#d4dee6;--warning:#84541b;--radius:8px}*{box-sizing:border-box}html{scroll-padding-top:90px}body{margin:0;color:var(--ink);background:var(--wash);font:16px/1.62 system-ui,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif}a{color:var(--accent);text-decoration-thickness:1px;text-underline-offset:3px}a:focus-visible,button:focus-visible,summary:focus-visible,input:focus-visible,select:focus-visible{outline:3px solid #d29a37;outline-offset:3px}button,input,select{font:inherit}button,.button{display:inline-block;border:1px solid var(--line);border-radius:5px;background:white;color:var(--ink);padding:.5rem .8rem;cursor:pointer;text-decoration:none}button:hover,.button:hover{border-color:var(--accent);background:#edf5f5}input,select{padding:.55rem;border:1px solid #91a4b4;border-radius:4px;max-width:100%;background:white;color:var(--ink)}header.site{background:var(--navy);color:white;padding:.85rem max(24px,calc((100% - 1440px)/2));display:flex;flex-wrap:wrap;align-items:center;gap:1rem}header.site a{color:#e8f5f5}header.site strong{margin-right:auto;letter-spacing:.03em}header.site nav{display:flex;flex-wrap:wrap;gap:1rem;font-size:.9rem}.layout{max-width:1440px;margin:auto;display:grid;grid-template-columns:245px minmax(0,1fr);gap:38px;padding:28px}.sidebar{position:sticky;top:20px;align-self:start;font-size:.9rem;max-height:92vh;overflow:auto}.sidebar a{display:block;padding:.32rem 0;text-decoration:none}.sidebar .sub{padding-left:12px;font-size:.84rem}.sidebar strong{display:block;margin:1rem 0 .35rem}.sidebar p{font-size:.85rem;color:var(--muted)}main{min-width:0}.hero{padding:15px 0 25px;border-bottom:2px solid var(--navy)}.eyebrow{font-size:.79rem;text-transform:uppercase;letter-spacing:.1em;font-weight:700;color:var(--accent)}h1{font-size:clamp(2rem,3.4vw,3rem);line-height:1.12;letter-spacing:-.035em;margin:.6rem 0 1rem}h2{font-size:1.65rem;line-height:1.25;margin:0 0 1rem;letter-spacing:-.02em}h3{font-size:1.22rem;line-height:1.35;margin:.2rem 0 .7rem}h4{font-size:.88rem;line-height:1.35;margin:0 0 .3rem;color:var(--muted)}p{margin:.6rem 0 1rem}.subtitle{font-size:1.25rem;color:var(--muted)}.lead{font-size:1.08rem}.actions{display:flex;gap:.6rem;flex-wrap:wrap;margin:1rem 0}.boundary,.note{padding:16px 20px;border-left:4px solid var(--accent);background:#eaf2f3;margin:1.2rem 0}.warning{border-left-color:var(--warning);background:#fcf5e8}.section{margin:40px 0}.group{margin:30px 0}.section-intro{color:var(--muted);max-width:75ch}.record{background:white;border:1px solid var(--line);border-radius:var(--radius);padding:24px;margin:18px 0;scroll-margin-top:20px}.record .operation{font-size:1.05rem;max-width:86ch}.id{font:700 .82rem ui-monospace,SFMono-Regular,monospace;color:var(--accent);margin-right:10px}.badges{display:flex;gap:.5rem;flex-wrap:wrap;margin:0 0 .8rem}.badge{font-size:.77rem;padding:.15rem .45rem;border:1px solid var(--line);border-radius:4px;color:var(--muted);background:var(--wash)}.badge.preview{color:var(--warning);background:#fcf5e8;border-color:#e9d0aa}.two{display:grid;grid-template-columns:1fr 1fr;gap:22px}.field{min-width:0}.field p{margin:.25rem 0 .9rem}.field.check{border-left:3px solid var(--accent);padding-left:14px;margin-top:16px}details{border-top:1px solid var(--line);margin-top:12px;padding-top:12px}summary{font-weight:650;cursor:pointer;line-height:1.4;margin-bottom:10px}.evidence{font-size:.8rem;line-height:1.5;color:var(--muted);padding-top:12px;margin-top:12px;border-top:1px solid var(--line);overflow-wrap:anywhere}.evidence a{display:inline}.toolbar{display:flex;gap:12px;align-items:end;flex-wrap:wrap;padding:18px;background:#e7edf1;border-radius:var(--radius)}.toolbar label{display:grid;gap:5px;font-size:.8rem;font-weight:650}.toolbar label:first-child{flex:1;min-width:180px}.toolbar input{width:100%}.result-count{font-size:.88rem;color:var(--muted)}.function-index{display:grid;grid-template-columns:1fr 1fr;gap:12px}.index-group{border:1px solid var(--line);border-radius:6px;padding:15px;background:white}.index-group a{display:block;font-size:.9rem;padding:.25rem 0}.index-group h3{font-size:1rem}.table-wrap{overflow-x:auto;border:1px solid var(--line);border-radius:6px;max-width:100%}table{border-collapse:collapse;width:100%;font-size:.92rem;background:white}th{text-align:left;background:var(--navy);color:white;font-weight:650}td,th{padding:12px 14px;border-bottom:1px solid var(--line);vertical-align:top;overflow-wrap:anywhere}tbody tr:nth-child(even){background:#f4f7f9}.steps{counter-reset:steps;list-style:none;padding:0}.steps li{counter-increment:steps;padding:0 0 16px 44px;position:relative}.steps li:before{content:counter(steps);position:absolute;left:0;top:1px;width:28px;height:28px;border-radius:50%;background:var(--navy);color:white;text-align:center;line-height:28px;font-size:.85rem}.steps strong{display:block}.steps .actor{font-size:.75rem;color:var(--accent);font-weight:700;text-transform:uppercase;letter-spacing:.05em}.steps p{margin:.25rem 0}.inputs{padding-left:22px}.inputs li{margin:.45rem 0}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#eff3f5;padding:16px;font: .87rem/1.55 ui-monospace,monospace;border-radius:5px}.sources{padding-left:0;list-style:none}.sources li{padding:12px 0;border-bottom:1px solid var(--line);overflow-wrap:anywhere}.sources small{display:block;color:var(--muted)}footer{border-top:1px solid var(--line);padding:24px;margin-top:36px;color:var(--muted);font-size:.85rem}.standalone{max-width:1150px;margin:0 auto;padding:36px 24px}.product-grid{display:grid;grid-template-columns:repeat(3,1fr);gap:22px}.product-grid article{background:white;border:1px solid var(--line);padding:25px;border-radius:8px}.product-grid article h2{font-size:1.45rem}.product-grid .scope-line{font-size:.93rem}.small{font-size:.85rem;color:var(--muted)}[hidden]{display:none!important}.skip{position:absolute;left:-9999px}.skip:focus{left:10px;top:10px;background:white;padding:12px;z-index:99}.presenter{max-width:1320px;margin:auto;padding:20px}.deck-controls{display:flex;gap:9px;align-items:center;flex-wrap:wrap;margin:12px 0}.deck-controls .counter{margin-left:auto;font-size:.9rem}.slide{background:white;border:1px solid var(--line);border-radius:8px;min-height:660px;padding:35px 42px 20px;margin:18px 0;position:relative}.slide h2{font-size:clamp(1.7rem,2.5vw,2.4rem);line-height:1.15;max-width:1000px}.slide .thesis{font-size:1.15rem;line-height:1.4;color:var(--accent);margin:0 0 24px}.slide table{font-size:1.08rem}.slide td,.slide th{padding:15px}.slide-foot{display:flex;flex-wrap:wrap;gap:10px;justify-content:space-between;margin-top:25px;padding-top:12px;border-top:1px solid var(--line);font-size:.8rem;color:var(--muted)}.flow{display:grid;gap:14px;grid-template-columns:repeat(var(--n),1fr);margin:35px 0}.flow>div{background:#eef4f5;padding:22px 18px;border-top:4px solid var(--accent);min-height:220px}.flow h3{font-size:1.2rem}.flow p{font-size:1.12rem;line-height:1.5}.flow .step-no{font-size:.8rem;color:var(--accent);font-weight:700}.speaker-notes{margin-top:15px;padding:18px;background:#eef3f6;border:1px solid var(--line);line-height:1.6}.demo-list{list-style:none;padding:0}.demo-list li{display:grid;grid-template-columns:80px 1fr;gap:12px;border-bottom:1px solid var(--line);padding:12px 0;font-size:1.05rem}.demo-list b{color:var(--accent)}.demo-safe{font-size:.92rem;padding:12px 15px;background:#fcf5e8;margin:15px 0}.timer{font-variant-numeric:tabular-nums}.deck-index{display:flex;flex-wrap:wrap;gap:5px;margin:12px 0}.deck-index button{font-size:.8rem;padding:.3rem .55rem}.deck-index button[aria-current=true]{background:var(--navy);color:white}body.present-mode header.site{display:none}body.present-mode .presenter{max-width:1450px}.progress-track{height:4px;background:#d6e2e7}.progress-track span{display:block;height:100%;background:var(--accent)}
@media(max-width:950px){.layout{grid-template-columns:1fr;padding:20px;gap:16px}.sidebar{position:static;max-height:none;border-bottom:1px solid var(--line);padding-bottom:15px}.sidebar nav{display:flex;flex-wrap:wrap;gap:8px 20px}.sidebar .sub,.sidebar p,.sidebar strong{display:none}.sidebar a{font-size:.85rem}.product-grid{grid-template-columns:1fr}.slide{min-height:0;padding:25px}.flow{grid-template-columns:1fr 1fr}.slide table{font-size:.95rem}.slide td,.slide th{padding:12px}.slide h2{font-size:1.8rem}}
@media(max-width:600px){body{font-size:15px}.two,.function-index,.flow{grid-template-columns:1fr}.layout{padding:16px}.record{padding:18px}.hero h1{font-size:2.1rem}.toolbar{align-items:stretch}.toolbar label{width:100%}.toolbar button{width:100%}header.site{padding:12px 16px}header.site nav{font-size:.8rem}.presenter{padding:10px}.slide{padding:21px 16px}.slide h2{font-size:1.6rem}.slide .thesis{font-size:1rem}.slide .table-wrap{overflow-x:auto}.slide table{min-width:580px}.slide .flow>div{min-height:0}.standalone{padding:26px 16px}.deck-controls .counter{margin-left:0}.demo-list li{grid-template-columns:55px 1fr}.evidence{font-size:.78rem}}
@media print{body{background:white;font-size:10pt}.site,.sidebar,.toolbar,.actions,.deck-controls,.deck-index,.progress-track,.no-print{display:none!important}.layout,.standalone,.presenter{display:block;max-width:none;padding:0;margin:0}.hero{padding:0}.record{break-inside:avoid;padding:15px;margin:12px 0}.section{margin:20px 0}.table-wrap{overflow:visible}.record[hidden],.group[hidden]{display:block!important}details{display:block}.slide{display:block!important;break-after:page;border:0;min-height:0;margin:0;padding:15px}.slide table{min-width:0;font-size:11pt}.slide h2{font-size:24pt}.slide .thesis{font-size:13pt}.slide .flow{grid-template-columns:repeat(var(--n),1fr)}.speaker-notes{display:none!important}.evidence a{color:#182b40}.slide-foot{font-size:9pt}a{color:#182b40}.product-grid{display:block}.product-grid article{margin-bottom:15px}.sources li{break-inside:avoid}}
'''

REPORT_JS='''
(()=>{const q=document.querySelector('#function-search'),g=document.querySelector('#function-group'),a=document.querySelector('#function-availability');
function filter(){const needle=(q?.value||'').toLowerCase().trim();let n=0;document.querySelectorAll('.function').forEach(x=>{const show=(!needle||x.textContent.toLowerCase().includes(needle))&&(!g.value||x.dataset.group===g.value)&&(!a.value||x.dataset.availability===a.value);x.hidden=!show;if(show)n++;});document.querySelectorAll('.function-group').forEach(x=>x.hidden=!x.querySelector('.function:not([hidden])'));const count=document.querySelector('#result-count');if(count)count.textContent=n+' functions shown';}
[q,g,a].forEach(x=>x?.addEventListener('input',filter));document.querySelector('#reset-filter')?.addEventListener('click',()=>{q.value='';g.value='';a.value='';filter();});
function showHash(){const id=decodeURIComponent(location.hash.slice(1)),target=document.getElementById(id);if(!target)return;const f=target.closest('.function');if(f&&f.hidden){q.value='';g.value='';a.value='';filter();}let node=target;while(node){if(node.tagName==='DETAILS')node.open=true;node=node.parentElement;}}window.addEventListener('hashchange',showHash);showHash();
document.querySelector('#expand-all')?.addEventListener('click',()=>document.querySelectorAll('main details').forEach(d=>d.open=true));document.querySelector('#collapse-all')?.addEventListener('click',()=>document.querySelectorAll('main details').forEach(d=>d.open=false));
let saved=[];window.addEventListener('beforeprint',()=>{saved=[...document.querySelectorAll('details')].map(d=>[d,d.open]);saved.forEach(([d])=>d.open=true)});window.addEventListener('afterprint',()=>saved.forEach(([d,o])=>d.open=o));document.querySelectorAll('[data-print]').forEach(x=>x.addEventListener('click',()=>window.print()));})();
'''

DECK_JS='''
(()=>{const slides=[...document.querySelectorAll('.slide')],buttons=[...document.querySelectorAll('.deck-index button')];let i=0,all=false,notes=false,elapsed=0,running=false,last=0;
function show(){slides.forEach((s,n)=>s.hidden=!all&&n!==i);buttons.forEach((b,n)=>b.setAttribute('aria-current',String(n===i)));document.querySelector('#counter').textContent=(i+1)+' / '+slides.length;document.querySelector('#progress').style.width=((i+1)/slides.length*100)+'%';document.querySelectorAll('.speaker-notes').forEach(n=>n.hidden=!notes);document.querySelector('#previous').disabled=i===0;document.querySelector('#next').disabled=i===slides.length-1;}
function move(n){i=Math.max(0,Math.min(slides.length-1,n));show();if(all)slides[i].scrollIntoView({block:'start'});}
document.querySelector('#previous').onclick=()=>move(i-1);document.querySelector('#next').onclick=()=>move(i+1);buttons.forEach((b,n)=>b.onclick=()=>move(n));
document.querySelector('#notes').onclick=()=>{notes=!notes;document.querySelector('#notes').setAttribute('aria-pressed',String(notes));show()};document.querySelector('#all-slides').onclick=()=>{all=!all;document.querySelector('#all-slides').setAttribute('aria-pressed',String(all));show()};document.querySelector('#present-mode').onclick=()=>document.body.classList.toggle('present-mode');
document.querySelector('#timer-toggle').onclick=()=>{running=!running;last=Date.now();document.querySelector('#timer-toggle').textContent=running?'Pause clock':'Start clock'};document.querySelector('#timer-reset').onclick=()=>{elapsed=0;last=Date.now();document.querySelector('#clock').textContent='00:00'};
setInterval(()=>{if(running){const now=Date.now();elapsed+=(now-last)/1000;last=now;}const t=Math.floor(elapsed);document.querySelector('#clock').textContent=String(Math.floor(t/60)).padStart(2,'0')+':'+String(t%60).padStart(2,'0')},250);
document.addEventListener('keydown',e=>{if(/INPUT|TEXTAREA|SELECT/.test(e.target.tagName))return;if(e.key==='ArrowRight'||e.key==='PageDown'){e.preventDefault();move(i+1)}if(e.key==='ArrowLeft'||e.key==='PageUp'){e.preventDefault();move(i-1)}if(e.key==='Home'){e.preventDefault();move(0)}if(e.key==='End'){e.preventDefault();move(slides.length-1)}if(e.key.toLowerCase()==='n'){notes=!notes;show()}if(e.key==='Escape')document.body.classList.remove('present-mode')});document.querySelector('[data-print]').onclick=()=>window.print();show();})();
'''

def shell(title,body,script='',prefix='',deck=False):
    return f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="description" content="Source-linked technical reference for AWS operations, application security and cloud cost teams."><title>{E(title)}</title><style>{CSS}</style></head><body><a class="skip" href="#main">Skip to content</a><header class="site"><strong>FRONTIER AGENTS / ENGINEERING REFERENCE</strong><nav><a href="{prefix}index.html">Collection</a><a href="{prefix}comparison.html">Platform comparison</a><a href="{prefix}METHOD.md">Method and limitations</a></nav></header>{body}<script>{script}</script></body></html>'

def source_register(d):
    out='<section class="section" id="sources"><h2>Official documentation</h2><p class="section-intro">Sources supporting the descriptions above. Review date: '+E(d['reviewed'])+'. Availability and UI paths can change. Proposed checks and procedures are editorial guidance, not vendor promises.</p><ol class="sources">'
    for key,s in d['sources'].items():
        out+=f'<li id="source-{E(key)}"><a href="{E(s["url"])}" target="_blank" rel="noopener noreferrer">{E(s["title"])}</a><small>{E(s["url"])}</small></li>'
    return out+'</ol></section>'

def render_report(d,deck):
    key=d['id'];total=sum(s['seconds'] for s in deck['slides']);demo=sum(x['seconds'] for x in deck['slides'] if x['kind']=='demo')
    nav='<aside class="sidebar"><strong>Read by task</strong><nav>'+''.join(f'<a href="#{anchor}">{label}</a>' for anchor,label in [('overview','Scope and object model'),('function-map','Function map'),('functions','Function reference'),('mechanisms','Configuration mechanisms'),('procedures','Operational procedures'),('gaps','Documentation qualifications'),('sources','Official sources')])+'</nav><strong>Function groups</strong>'+''.join(f'<a class="sub" href="#group-{g["id"]}">{E(g["title"])}</a>' for g in d['groups'])+p('Functions describe product behavior. Mechanisms explain configuration. Procedures describe a proposed engineering decision. No live product execution is claimed.')+'</aside>'
    body='<div class="layout">'+nav+'<main id="main"><section class="hero" id="overview"><div class="eyebrow">Technical reference · reviewed '+E(d['reviewed'])+'</div><h1>'+E(d['title'])+'</h1>'+p(d['subtitle'],'subtitle')+p(d['scope'],'lead')
    body+=f'<div class="actions"><a class="button" href="presentation.html">Present · {stamp(total)}</a><a class="button" href="../presentations/{PPT[key]}">Editable PowerPoint</a><a class="button" href="report.md">Markdown reference</a><button data-print>Print reference</button></div><p class="small">Presentation includes {stamp(demo)} of prepared-result demonstration. Preparation, service processing time and questions are additional.</p>'
    body+='<div class="boundary warning"><strong>Availability and scope</strong>'+p(d['availability'])+refs(d,d['availability_refs'])+'</div></section>'
    body+='<section class="section"><h2>Understand the product objects</h2>'+table(['Object','Meaning'],[[x[0],x[1]] for x in d['objects']])+refs(d,[x[2] for x in d['objects']])+'</section>'
    body+='<section class="section" id="function-map"><h2>Function map</h2><p class="section-intro">Each entry names an operation and links to its input, behavior, result and limitations. Groups organize responsibility; they are not additional capabilities.</p><div class="function-index">'
    for g in d['groups']:
        body+='<div class="index-group"><h3>'+E(g['title'])+'</h3>'+p(g['purpose'],'small')+''.join(f'<a href="#{f["id"]}"><span class="id">{f["id"]}</span>{E(f["name"])}</a>' for f in d['functions'] if f['group']==g['id'])+'</div>'
    body+='</div></section><section class="section" id="functions"><h2>Function reference</h2><div class="toolbar"><label>Search function details<input id="function-search" type="search" placeholder="For example: review, permissions, evidence"></label><label>Group<select id="function-group"><option value="">All groups</option>'+''.join(f'<option value="{g["id"]}">{E(g["title"])}</option>' for g in d['groups'])+'</select></label><label>Availability<select id="function-availability"><option value="">All documented functions</option><option value="Preview">Preview</option><option value="GA">Explicitly GA</option><option value="Documented">Documented; see feature qualification</option></select></label><button id="reset-filter">Reset</button></div><p id="result-count" class="result-count" aria-live="polite">'+str(len(d['functions']))+' functions shown</p><p class="small">“Documented” is an evidence label, not an assertion of general availability. Activation requirements are listed separately.</p>'
    for g in d['groups']:
        body+=f'<section class="group function-group" id="group-{g["id"]}"><h3>{E(g["title"])}</h3>{p(g["purpose"],"section-intro")}'
        for f in [x for x in d['functions'] if x['group']==g['id']]:
            badge='preview' if f['availability']=='Preview' else ''
            activation_badge='' if f['activation'].startswith('See invocation') else '<span class="badge">'+E(f['activation'])+'</span>'
            body+=f'<article class="record function" id="{f["id"]}" data-group="{g["id"]}" data-availability="{f["availability"]}"><div class="badges"><span class="id">{f["id"]}</span><span class="badge {badge}">{E(f["availability"])}</span>{activation_badge}</div><h3>{E(f["name"])}</h3>{p(f["operation"],"operation")}<div class="two">{block("Input",f["input"])}{block("Output and interpretation",f["output"])}</div><details><summary>Invocation, limitations and proposed verification</summary>{block("Where to invoke / prerequisite",f["entry"])}{block("Limitations and effects",f["limits"])}{block("Proposed operator verification — not an executed test",f["check"],"check")}</details>{refs(d,f["refs"])}</article>'
        body+='</section>'
    body+='</section><section class="section" id="mechanisms"><h2>Configuration mechanisms</h2><p class="section-intro">How to enable or constrain the functions above. Configuration details are not counted as separate use cases. Recommended verification exercises have not been run against an AWS account.</p><div class="actions"><button id="expand-all">Expand technical detail</button><button id="collapse-all">Collapse technical detail</button></div>'
    for m in d['mechanisms']:
        body+=f'<article class="record" id="{m["id"]}"><span class="id">{m["id"]}</span><h3>{E(m["name"])}</h3>{p(m["explanation"],"operation")}<h4>Configuration sequence</h4><ol class="steps">'+''.join('<li>'+E(step)+'</li>' for step in m['configure'])+'</ol>'+block('Lifecycle and update behavior',m['lifecycle'])+block('Proposed verification — not an executed test',m['verification'],'check')+refs(d,m['refs'])+'</article>'
    body+='</section><section class="section" id="procedures"><h2>Operational procedures</h2><p class="section-intro">Proposed procedures for a specific engineering decision. They compose named functions and include stop conditions. They are not stories of successful customer deployments.</p>'
    for u in d['procedures']:
        body+=f'<article class="record procedure" id="{u["id"]}"><span class="id">{u["id"]}</span><span class="badge">Proposed procedure</span><h3>{E(u["title"])}</h3>{p(u["goal"],"lead")}<div class="two">{block("Use when",u["when"])}{block("Do not use this to claim",u["not_for"])}</div><p class="small">Uses: '+', '.join(f'<a href="#{x}">{x}</a>' for x in u['functions'])+'</p><h4>Required inputs</h4><ul class="inputs">'+''.join('<li>'+E(x)+'</li>' for x in u['inputs'])+'</ul><h4>Execution and decision points</h4><ol class="steps">'+''.join(f'<li><span class="actor">{E(actor)}</span><strong>{E(title)}</strong>{p(desc)}</li>' for actor,title,desc in u['steps'])+'</ol><div class="two">'+block('Acceptance evidence',u['acceptance'])+block('Stop / escalate',u['stop'])+'</div><details><summary>Example instruction and demonstration runbook</summary><p class="small">Example instruction; adapt identifiers and authorization to the actual environment.</p><pre>'+E(u['prompt'])+'</pre>'
        if u['demo_minutes']:
            body+=f'<h4>{u["demo_minutes"]}-minute prepared-result walkthrough</h4>'+table(['Budget','Presenter action'],[[f'{mins} min',step] for mins,step in u['demo']])
        else:body+=p('Reference procedure only; no demonstration allocation in the core presentation.','small')
        body+=block('Fallback',u['fallback'])+'</details>'+refs(d,u['refs'])+'</article>'
    body+='</section>'
    if 'worked_example' in d:
        w=d['worked_example'];body+='<section class="section" id="worked-example"><h2>'+E(w['title'])+'</h2><div class="boundary warning">'+E(w['warning'])+'</div>'+table(w['headers'],w['rows'])+p(w['explanation'])+'</section>'
    body+='<section class="section" id="gaps"><h2>Documentation qualifications</h2>'+''.join('<div class="note">'+E(x)+'</div>' for x in d['gaps'])+'</section>'+source_register(d)+'<footer>No live AWS account, assessment, billing query or memory-recall experiment was executed for this reference. No accuracy, productivity, security coverage or savings outcome is measured.</footer></main></div>'
    return shell(d['title']+' — technical reference',body,REPORT_JS,'../')

def slide_body(d,s):
    if s['kind']=='demo':
        u=next(x for x in d['procedures'] if x['id']==s['procedure'])
        return '<ol class="demo-list">'+''.join(f'<li><b>{mins} min</b><span>{E(step)}</span></li>' for mins,step in u['demo'])+'</ol><div class="demo-safe"><strong>Fallback / safety:</strong> '+E(u['fallback'])+'</div>'
    if s['kind'] in ('flow','cover'):
        return '<div class="flow" style="--n:'+str(len(s['headers']))+'">'+''.join(f'<div><span class="step-no">{n+1:02d}</span><h3>{E(h)}</h3>{p(s["rows"][0][n])}</div>' for n,h in enumerate(s['headers']))+'</div>'
    return table(s['headers'],s['rows'])

def render_deck(d,deck):
    slides=deck['slides'];total=sum(s['seconds'] for s in slides);demo=sum(s['seconds'] for s in slides if s['kind']=='demo')
    out=f'<main id="main" class="presenter"><p class="eyebrow">{E(d["title"])} · presenter view</p><p class="small">{stamp(total)} total, including {stamp(demo)} prepared-result demonstrations. Service execution time, preparation and questions are additional. <a href="{REPORT[d["id"]]}">Full technical reference</a> · <a href="presentation-transcript.md">Transcript</a> · <a href="../presentations/{PPT[d["id"]]}">PowerPoint</a></p><div class="deck-controls"><button id="previous">Previous</button><button id="next">Next</button><button id="notes" aria-pressed="false">Speaker notes (N)</button><button id="all-slides" aria-pressed="false">All slides</button><button id="present-mode">Focus view</button><button data-print>Print slides</button><span id="counter" class="counter" aria-live="polite"></span></div><div class="deck-controls"><button id="timer-toggle">Start clock</button><button id="timer-reset">Reset clock</button><span id="clock" class="timer">00:00</span><span class="small">Elapsed session clock; no automatic slide advance.</span></div><div class="deck-index">'+''.join(f'<button aria-label="Slide {i+1}: {E(s["title"])}">{i+1}</button>' for i,s in enumerate(slides))+'</div><div class="progress-track"><span id="progress"></span></div>'
    elapsed=0
    for i,s in enumerate(slides):
        end=elapsed+s['seconds'];out+=f'<section class="slide" id="slide-{i+1}"><div class="eyebrow">{E(d["title"])} / {i+1:02d}</div><h2>{E(s["title"])}</h2>{p(s["thesis"],"thesis")}'+slide_body(d,s)
        if s.get('callout'):out+='<div class="boundary">'+E(s['callout'])+'</div>'
        out+=f'<div class="slide-foot"><span>{stamp(elapsed)}–{stamp(end)} · {stamp(s["seconds"])} allocation</span><span>{"Prepared-result demonstration" if s["kind"]=="demo" else "Explanation and discussion"}</span><span>{i+1} / {len(slides)}</span></div>'+refs(d,s['refs'],REPORT[d['id']])+f'<aside class="speaker-notes" hidden><strong>Presenter notes</strong>{p(s["notes"])}</aside></section>';elapsed=end
    return shell(d['title']+' — presentation',out+'</main>',DECK_JS,'../',True)

def md_refs(d,ids):return 'Official documentation: '+ ' · '.join(f'[{d["sources"][i]["title"]}]({d["sources"][i]["url"]})' for i in dict.fromkeys(ids))+'\n'
def render_markdown(d):
    lines=[f'# {d["title"]}',d['subtitle'],f'Reviewed: {d["reviewed"]}',d['scope'],'## Availability',d['availability'],md_refs(d,d['availability_refs']),'## Product objects']
    for name,definition,ref in d['objects']:lines.extend([f'### {name}',definition,md_refs(d,[ref])])
    lines.append('## Product functions')
    for g in d['groups']:
        lines.extend([f'### {g["title"]}',g['purpose']])
        for f in d['functions']:
            if f['group']!=g['id']:continue
            lines.extend([f'#### {f["id"]} — {f["name"]}',f'Availability: {f["availability"]}. Activation: {f["activation"]}.',f['operation']])
            for k,label in [('input','Input'),('output','Output'),('entry','Invocation and prerequisites'),('limits','Limitations and effects'),('check','Proposed operator verification — not executed')]:lines.extend([f'**{label}.** {f[k]}'])
            lines.append(md_refs(d,f['refs']))
    lines.append('## Configuration mechanisms')
    for m in d['mechanisms']:
        lines.extend([f'### {m["id"]} — {m["name"]}',m['explanation'],'\n'.join(f'{i+1}. {x}' for i,x in enumerate(m['configure'])),f'**Lifecycle.** {m["lifecycle"]}',f'**Proposed verification.** {m["verification"]}',md_refs(d,m['refs'])])
    lines.append('## Proposed operational procedures')
    for u in d['procedures']:
        lines.extend([f'### {u["id"]} — {u["title"]}',u['goal'],f'**Use when:** {u["when"]}',f'**Not for:** {u["not_for"]}',f'**Functions used:** {", ".join(u["functions"])}','**Inputs**','\n'.join('- '+x for x in u['inputs'])])
        for i,(actor,title,desc) in enumerate(u['steps']):lines.append(f'{i+1}. **{actor} — {title}.** {desc}')
        lines.extend([f'**Acceptance:** {u["acceptance"]}',f'**Stop:** {u["stop"]}','**Example instruction**','```text\n'+u['prompt']+'\n```',f'**Demonstration allocation:** {u["demo_minutes"]} minutes.','\n'.join(f'- {n} min: {step}' for n,step in u['demo']),f'**Fallback:** {u["fallback"]}',md_refs(d,u['refs'])])
    if 'worked_example' in d:
        w=d['worked_example'];lines.extend(['## '+w['title'],w['warning'],table_md(w['headers'],w['rows']),w['explanation']])
    lines.extend(['## Documentation qualifications',*d['gaps'],'## Official source register'])
    lines.extend(f'- [{s["title"]}]({s["url"]})' for s in d['sources'].values())
    lines.append('No live product execution or measured operational benefit is claimed.')
    return '\n\n'.join(lines)+'\n'

def table_md(headers,rows):
    return '| '+' | '.join(map(str,headers))+' |\n|'+'|'.join('---' for _ in headers)+'|\n'+'\n'.join('| '+' | '.join(str(v).replace('|','/') for v in row)+' |' for row in rows)

def render_transcript(d,deck):
    total=sum(s['seconds'] for s in deck['slides']);demo=sum(s['seconds'] for s in deck['slides'] if s['kind']=='demo');lines=[f'# {d["title"]}: presentation transcript',f'Total {stamp(total)}, including {stamp(demo)} prepared-result demonstrations. No service-completion-time guarantee. Questions, preparation and live setup are additional.'];elapsed=0
    for i,s in enumerate(deck['slides']):
        end=elapsed+s['seconds'];lines.extend([f'## {i+1}. {s["title"]}',f'{stamp(elapsed)}–{stamp(end)}; allocation {stamp(s["seconds"])}',s['thesis']])
        if s['kind']=='demo':
            u=next(u for u in d['procedures'] if u['id']==s['procedure']);lines.extend([table_md(['Minutes','Action'],u['demo']),f'Fallback: {u["fallback"]}'])
        else:lines.append(table_md(s['headers'],s['rows']))
        if s.get('callout'):lines.append(s['callout'])
        lines.extend(['### Presenter notes',s['notes']])
        if s['refs']:lines.append(md_refs(d,s['refs']))
        elapsed=end
    return '\n\n'.join(lines)+'\n'

NAVY='11283E';INK='182B40';TEAL='006D76';MUTED='536477';WHITE='FFFFFF';PALE='EDF3F5';LINE='D4DEE6'
def shape(slide,x,y,w,h,fill,line=None,kind=MSO_SHAPE.RECTANGLE):
    sh=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background() if line is None else None
    if line:sh.line.color.rgb=RGBColor.from_string(line);sh.line.width=Pt(.7)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh

def text(slide,value,x,y,w,h,size=18,color=INK,bold=False):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    for i,line in enumerate(str(value).split('\n')):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.text=line;para.font.name='Arial';para.font.size=Pt(size);para.font.bold=bold;para.font.color.rgb=RGBColor.from_string(color);para.space_after=Pt(4);para.line_spacing=1.08
    return sh

def pptx_deck(d,deck):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5);prs.core_properties.title=d['title']+' — engineering presentation';prs.core_properties.subject=d['subtitle'];prs.core_properties.author='Engineering documentation';elapsed=0
    for i,s in enumerate(deck['slides']):
        sl=prs.slides.add_slide(prs.slide_layouts[6]);shape(sl,0,0,13.333,7.5,WHITE);shape(sl,0,0,13.333,.13,TEAL)
        text(sl,d['title'].upper()+'  /  ENGINEERING REFERENCE',.5,.32,11.8,.25,10,TEAL,True)
        text(sl,s['title'],.5,.78,12.25,1.0,29,INK,True)
        text(sl,s['thesis'],.5,1.94,12.15,.52,17,TEAL)
        top=2.55;bottom=6.65
        if s['kind'] in ('flow','cover'):
            count=len(s['headers']);gap=.18;w=(12.3-gap*(count-1))/count
            for j,(heading,body) in enumerate(zip(s['headers'],s['rows'][0])):
                x=.5+j*(w+gap);shape(sl,x,top+.2,w,2.95,PALE);shape(sl,x,top+.2,w,.07,TEAL);text(sl,f'{j+1:02d}',x+.18,top+.42,w-.36,.3,12,TEAL,True);text(sl,heading,x+.18,top+.9,w-.36,.7,21,INK,True);text(sl,body,x+.18,top+1.67,w-.36,1.12,18,INK)
        elif s['kind']=='table':
            n=len(s['headers']);widths=([3.3,4.5,4.5] if n==3 else [2.3,3.34,3.33,3.33] if n==4 else [12.3/n]*n);rows=s['rows'];height=3.6 if s.get('callout') else 4.08;hh=.55;rh=(height-hh)/len(rows)
            x=.5
            for j,head in enumerate(s['headers']):shape(sl,x,top,widths[j],hh,NAVY);text(sl,head,x+.14,top+.12,widths[j]-.28,hh-.16,15,WHITE,True);x+=widths[j]
            for r,row in enumerate(rows):
                x=.5;y=top+hh+r*rh
                for c,val in enumerate(row):
                    shape(sl,x,y,widths[c],rh,WHITE if r%2 else PALE,LINE);text(sl,val,x+.14,y+.13,widths[c]-.28,rh-.17,17 if len(str(val))<130 else 16,INK,c==0);x+=widths[c]
        elif s['kind']=='demo':
            u=next(x for x in d['procedures'] if x['id']==s['procedure']);rh=3.1/len(u['demo'])
            for r,(mins,desc) in enumerate(u['demo']):
                y=top+r*rh;shape(sl,.5,y,12.3,rh-.04,PALE);text(sl,f'{mins} min',.68,y+.12,.8,.35,17,TEAL,True);text(sl,desc,1.65,y+.1,10.9,rh-.12,18,INK)
            shape(sl,.5,5.86,12.3,.82,'FCF5E8');text(sl,'PREPARED-EVIDENCE FALLBACK',.68,5.98,11.9,.22,10,'84541B',True);text(sl,u['fallback'],.68,6.25,11.9,.37,13,INK)
        if s.get('callout'):
            shape(sl,.5,6.25,12.3,.47,NAVY);text(sl,s['callout'],.65,6.34,12,.3,12,WHITE,True)
        end=elapsed+s['seconds'];shape(sl,.5,6.95,12.3,.012,LINE);text(sl,f'{stamp(elapsed)}–{stamp(end)}   |   {stamp(s["seconds"])} allocation',.5,7.1,5,.22,10,MUTED);text(sl,f'{i+1:02d} / {len(deck["slides"]):02d}',11.85,7.1,.95,.22,10,MUTED)
        refline='Sources: '+', '.join(s['refs'])+' · full URLs in notes and report' if s['refs'] else 'Synthetic illustration — not measured service output'
        text(sl,refline,5.25,7.1,6.4,.22,9,MUTED)
        note=f'{s["title"]}\nTime: {stamp(elapsed)}–{stamp(end)}; budget {stamp(s["seconds"])}.\n\n{s["notes"]}\n\n'
        if s['kind']=='demo':note+='\n'.join(f'{mins} min — {desc}' for mins,desc in u['demo'])+'\nFallback: '+u['fallback']+'\n\n'
        note+='OFFICIAL SOURCES\n'+'\n'.join(d['sources'][ref]['title']+'\n'+d['sources'][ref]['url'] for ref in s['refs']);sl.notes_slide.notes_text_frame.text=note;elapsed=end
    out=ROOT/'presentations'/PPT[d['id']];out.parent.mkdir(exist_ok=True);prs.save(out)

def compare_page(docs):
    # Each cell links to the product-specific operation, not an invented shared implementation.
    rows=[
      ('Primary evidence',('devops','D05','Current telemetry, topology and change context for an incident.'),('security','S03','Design/source assessments and executed target-test evidence.'),('finops','F01','Billing-tool results and source optimization recommendations.')),
      ('Historical context',('devops','DM3','Managed monitor memories and corrective reflections; explicit store lifecycle.'),('security','SM4','Requirements, source context and linked retest history; no general incident-memory claim.'),('finops','FM2','User context files and retained reporting preferences; no changes to billing truth.')),
      ('Change proposal',('devops','D07','Mitigation proposal or prevention recommendation.'),('security','S08','Proposed code repair, followed by developer tests and deployment.'),('finops','F06','Source recommendation summarized for workload-owner review.')),
      ('Resource/application effects',('devops','DM5','Approved directed actions; preview Sandbox and release testing have separate boundaries.'),('security','S06','Authorized penetration-test traffic against a configured target.'),('finops','FM4','Documented external publishing, not resource resizing or commitment purchase.')),
      ('Recurring work',('devops','DM4','Custom-agent definition, selected tools/context and triggers.'),('security','SM2','Source-review trigger rules and assessment jobs; no universal scheduler assumed.'),('finops','FM3','Scheduled tasks or supported cost-anomaly event automations.')),
      ('Meaning of success',('devops','DU1','Diagnosis supported by current evidence; recovery verified separately.'),('security','SU2','Specific finding retested under valid conditions; regressions independently checked.'),('finops','FU1','Matched totals reconcile; explanation and estimates remain qualified.'))
    ]
    body='<main class="standalone" id="main"><div class="eyebrow">Cross-platform comparison</div><h1>Compare boundaries, not product slogans</h1><p class="lead">The three services do not share a documented permission model, memory store or execution runtime. This comparison links each distinction to its product reference. It is not a proposed cross-agent orchestration architecture.</p><div class="table-wrap"><table><thead><tr><th>Dimension</th>'+''.join('<th>'+E(docs[k]['title'])+'</th>' for k in KEYS)+'</tr></thead><tbody>'
    mdrows=[]
    for row in rows:
        body+='<tr><th style="background:#edf3f5;color:#182b40">'+E(row[0])+'</th>';vals=[row[0]]
        for key,anchor,value in row[1:]:
            href=f'{key}/{REPORT[key]}#{anchor}';body+='<td>'+E(value)+f' <a href="{href}">{anchor}</a></td>';vals.append(f'{value} [{anchor}]({href})')
        body+='</tr>';mdrows.append(vals)
    body+='</tbody></table></div><section class="section"><h2>What a team can standardize</h2><p>Standardize the work record: product and run ID; source revision or data cutoff; resource or target scope; evidence links; proposed action; approval owner; independent verification; unresolved gaps. These are team process fields, not an assertion that all three products emit the same schema.</p><h2>What must stay product-specific</h2><p>Keep access grants, memory activation, source connections, test authentication, schedules, approval gates and deletion behavior specific to each service. A permission or successful result in one product must not be treated as authorization or proof in another.</p></section></main>'
    write('comparison.html',shell('Platform comparison',body));write('comparison.md','# Platform comparison\n\nNo shared runtime, memory or authorization model is asserted.\n\n'+table_md(['Dimension',*[docs[k]['title'] for k in KEYS]],mdrows)+'\n\nThe shared work-record fields are a proposed team practice, not a native common schema.\n')

def index_page(docs,teaching):
    body='<main class="standalone" id="main"><div class="eyebrow">Technical reference · 21 October 2026</div><h1>Functions, configuration and operating decisions</h1><p class="lead">Three independently documented AWS services. Select a product to inspect its actual operations, enablement mechanisms and proposed procedures. Presentation material is generated from the same reference, with separate time allocations for prepared-result demonstrations.</p><div class="boundary"><strong>Read the distinction first.</strong> A function describes what the product does. A mechanism describes how context, permissions or triggers enable it. A procedure composes functions to answer a specific engineering question. Proposed checks are not vendor guarantees or claims of executed tests.</div><div class="product-grid">'
    for key in KEYS:
        d=docs[key];deck=teaching[key];total=sum(s['seconds'] for s in deck['slides']);demo=sum(s['seconds'] for s in deck['slides'] if s['kind']=='demo');body+=f'<article><div class="eyebrow">{E(key)}</div><h2>{E(d["title"])}</h2>{p(d["subtitle"],"scope-line")}<p class="small">'+', '.join(E(g['title']) for g in d['groups'])+f'</p><p><a href="{key}/{REPORT[key]}">Read technical reference</a></p><p><a href="{key}/presentation.html">Present {stamp(total)}</a> · <a href="presentations/{PPT[key]}">PowerPoint</a></p><p class="small">Includes {stamp(demo)} of prepared-result demonstration. Questions and live setup are additional.</p></article>'
    body+='</div><section class="section"><h2>Use the reports as references, not assurance statements</h2><p>Each function identifies its inputs, operation, output, invocation and limitations. Configuration sections explain enabling and lifecycle behavior. Procedures include required evidence, decision points and stop conditions. The source register links to official documentation and records qualifications where pages or versions differ.</p><p><a href="comparison.html">Compare platform boundaries</a> · <a href="METHOD.md">Read the method and evidence boundary</a> · <a href="VALIDATION.md">Review validation results</a></p></section><footer>No live AWS workflow, scan, cost query, resource change or memory-recall experiment was executed. No operational benefit or service accuracy is measured. Git history holds superseded revisions; no manual archive is maintained.</footer></main>'
    write('index.html',shell('AWS agents — engineering reference',body))

def root_docs(docs,teaching):
    timing=[]
    for k in KEYS:
        total=sum(s['seconds'] for s in teaching[k]['slides']);demo=sum(s['seconds'] for s in teaching[k]['slides'] if s['kind']=='demo');timing.append([docs[k]['title'],f'[{stamp(total)} presentation]({k}/presentation.html)',stamp(demo),f'[Report]({k}/{REPORT[k]})',f'[PowerPoint](presentations/{PPT[k]})'])
    write('README.md','# AWS agents — engineering reference\n\nThis is a replacement of the earlier narrative-led collection, not a format patch. The current scope is AWS DevOps Agent, AWS Security Agent and AWS FinOps Agent.\n\nOpen `index.html` after cloning the repository. The reports and browser presentations are self-contained and require no server. GitHub previews display source; this is not a GitHub Pages deployment.\n\n'+table_md(['Product','Total allocation','Demo included','Reference','Editable deck'],timing)+'\n\n## Structure\n\n`reference/` holds the canonical product descriptions and teaching material. Each product directory contains a grouped function reference, configuration mechanisms, proposed procedures, browser presentation and Markdown transcript. `comparison.html` compares concrete product boundaries; it does not invent a common agent runtime.\n\nThe function catalogue is organized by product responsibility. Setup mechanisms are not counted as capabilities, and operational examples are not presented as executed results. Each material product assertion links to an official source; proposed checks, synthetic arithmetic and documentation qualifications are labeled separately. See [METHOD.md](METHOD.md).\n\n## Build and validate\n\n```sh\npython -m pip install python-pptx==1.0.2 beautifulsoup4==4.13.4 playwright==1.57.0\npython -m playwright install chromium\npython tools/build_reference.py\npython tools/check_reference.py\n```\n\nThe publication workflow validates before committing each changed file or deletion separately. Earlier revisions remain in Git history. No archive directory, distribution ZIP, encoded import package or historical copy is generated.\n\n## Evidence boundary\n\nSource review: 21 October 2026. Documents can change after that date. The references state what the official pages document, not an independently observed service implementation. No live product execution or measured accuracy, security coverage, cost saving or productivity benefit is claimed. Browser and presentation checks are recorded in [VALIDATION.md](VALIDATION.md); they do not establish factual or operational correctness.\n')
    write('METHOD.md','''# Method and evidence boundary

## The unit of documentation

A product function is an operation on an identifiable input that produces an inspectable result. A separate invocation, result type or effect boundary is a reason to split functions. Similar examples of the same operation are not separate functions. Related lifecycle controls can be documented together when they operate on the same record.

A configuration mechanism explains how access, context, instructions, memory, authentication or triggers support functions. A proposed operational procedure composes functions to make an engineering decision. It must identify when it applies, required evidence, responsible actors, decisions, stop conditions and completion evidence. There is no prescribed number of stages.

## Evidence labels

- **Documented product behavior:** supported by the official source links attached to the record. “Documented” does not itself mean generally available; feature availability and activation requirements are separate fields.
- **Proposed operator verification or procedure:** an editorial recommendation, not a promise of a built-in feature and not an account of a completed test.
- **Synthetic illustration:** invented values used to expose calculation or decision logic. They are not service prices, customer data or measured product output.
- **Documentation qualification:** a known source conflict, version difference or unestablished implementation detail. Prefer the specific current feature page and dated change history over an older general summary; preserve remaining uncertainty.

## Source use

Only official AWS product documentation and official availability announcements support product-specific claims. Each function and mechanism links to the relevant pages; each presentation links its source identifiers to the same register. The authored descriptions are paraphrases, not copies of vendor positioning. Review date is 21 October 2026. The documentation is not a specification of hidden model architecture or training.

## Presentation design

Presentations explain decisions and differences using editable tables, flows and worked examples. They do not reproduce the complete reference one card per slide. Full details remain in the linked report. Every timed slide includes its budget and cumulative interval; notes retain sources and the intended explanation. Demonstrations are prepared-result walkthroughs with fallbacks. Live job completion, setup and questions are outside those budgets.

## Validation is layered

Content review checks taxonomy, distinct inputs/results, source support, setup specificity and failure interpretations. Automated checks cover schema, cross-references, arithmetic, local links, browser controls, mobile layout, presentation bounds and notes. Rendering and visual inspection are separate from bounds checks. None replaces execution against the actual AWS services.

No live AWS account was queried or changed, no penetration test ran, and no cost or memory-recall result was measured. The validation record must not describe an unperformed test as passing.
''')
    write('SCOPE.md','# Scope\n\nThe current collection covers AWS DevOps Agent, AWS Security Agent and AWS FinOps Agent. Shared comparison material is limited to their documented evidence, configuration and action boundaries. Previous content is retained only by Git history. No manual archives are maintained.\n')

def build():
    docs={k:load(k) for k in KEYS};teaching=json.loads((ROOT/'reference/teaching.json').read_text(encoding='utf-8'))
    for key in KEYS:
        d=docs[key];deck=teaching[key];write(f'{key}/{REPORT[key]}',render_report(d,deck));write(f'{key}/report.md',render_markdown(d));write(f'{key}/presentation.html',render_deck(d,deck));write(f'{key}/presentation-transcript.md',render_transcript(d,deck));pptx_deck(d,deck)
    compare_page(docs);index_page(docs,teaching);root_docs(docs,teaching)
    write('qa/timing.json',json.dumps({k:{'seconds':sum(s['seconds'] for s in teaching[k]['slides']),'demo_seconds':sum(s['seconds'] for s in teaching[k]['slides'] if s['kind']=='demo'),'slides':len(teaching[k]['slides'])} for k in KEYS},indent=2)+'\n')
    print('Built technical references, presentations, transcripts and product comparison.')

def commit_changes():
    def git(*args):return subprocess.check_output(['git',*args],cwd=ROOT)
    names=set(git('diff','--name-only','-z').decode().split('\0'))|set(git('ls-files','--others','--exclude-standard','-z').decode().split('\0'))
    for name in sorted(names-{''}):
        git('add','-A','--',name)
        if subprocess.run(['git','diff','--cached','--quiet','--',name],cwd=ROOT).returncode:
            git('commit','-m',f'docs: replace {name}')
    print('Committed each changed file or deletion separately.')

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--commit',action='store_true');args=ap.parse_args();commit_changes() if args.commit else build()
