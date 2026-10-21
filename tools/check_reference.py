#!/usr/bin/env python3
"""Check generated output without claiming runtime product validation."""
from __future__ import annotations
import hashlib
import json
import os
from decimal import Decimal
from pathlib import Path
from urllib.parse import unquote, urlsplit
from bs4 import BeautifulSoup
from pptx import Presentation
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).resolve().parents[1]
KEYS=('devops','security','finops')
REPORT={'devops':'validated-report.html','security':'report.html','finops':'report.html'}
errors=[]
checks=[]

def navigate(page,path):
    if os.environ.get('HTML_LOAD_MODE')=='content':
        page.set_content(path.read_text(encoding='utf-8'),wait_until='load')
    else:
        page.goto(path.as_uri())

def require(condition,message):
    if not condition:errors.append(message)

def refs_ok(d,refs,label):
    require(bool(refs),label+' has no product evidence references')
    for ref in refs:require(ref in d['sources'],label+' unknown source '+str(ref))

teaching=json.loads((ROOT/'reference/teaching.json').read_text())
docs={k:json.loads((ROOT/'reference'/f'{k}.json').read_text()) for k in KEYS}
for key,d in docs.items():
    groups=[g['id'] for g in d['groups']];fids=[f['id'] for f in d['functions']];names=[f['name'].casefold() for f in d['functions']]
    require(len(fids)==len(set(fids)),key+' duplicate function ID');require(len(names)==len(set(names)),key+' duplicate function name')
    for src in d['sources'].values():require(urlsplit(src['url']).hostname in ('docs.aws.amazon.com','aws.amazon.com'),key+' nonofficial source '+src['url'])
    refs_ok(d,d['availability_refs'],key+' availability')
    for f in d['functions']:
        require(f['group'] in groups,key+' invalid group '+f['id'])
        require(f['availability'] in ('Documented','Preview','GA'),key+' conflated availability '+f['id'])
        for field in ('input','operation','output','entry','limits','check','activation'):
            require(isinstance(f.get(field),str) and len(f[field].strip())>10,key+' empty/short '+field+' '+f['id'])
        refs_ok(d,f['refs'],key+'/'+f['id'])
    for m in d['mechanisms']:
        require(len(m['configure'])>=3,key+' incomplete mechanism '+m['id']);refs_ok(d,m['refs'],key+'/'+m['id'])
    for u in d['procedures']:
        require(all(x in fids for x in u['functions']),key+' unknown function in '+u['id'])
        require(all(u.get(x) for x in ('when','not_for','inputs','steps','acceptance','stop','fallback')),key+' incomplete procedure '+u['id'])
        require(sum(x[0] for x in u['demo'])==u['demo_minutes'],key+' demo timing mismatch '+u['id']);refs_ok(d,u['refs'],key+'/'+u['id'])
    for s in teaching[key]['slides']:
        require(s['seconds']>0 and len(s['notes'])>100,key+' incomplete timed slide '+s['title'])
        for r in s['refs']:require(r in d['sources'],key+' unknown slide source '+r)
        if s['kind']=='demo':
            u=next((u for u in d['procedures'] if u['id']==s['procedure']),None);require(u is not None,key+' unknown demo procedure')
            if u:require(u['demo_minutes']*60==s['seconds'],key+' slide/procedure demo mismatch')
        else:require(all(len(row)==len(s['headers']) for row in s['rows']),key+' malformed slide table '+s['title'])
    checks.append({'product':key,'functions':len(fids),'groups':len(groups),'mechanisms':len(d['mechanisms']),'procedures':len(d['procedures']),'timed_seconds':sum(s['seconds'] for s in teaching[key]['slides']),'demo_seconds':sum(s['seconds'] for s in teaching[key]['slides'] if s['kind']=='demo')})

# Validate the pedagogical arithmetic without importing any pricing assertion.
w=docs['finops']['worked_example'];base=Decimal('0');current=Decimal('0')
for row in w['rows']:
    q0,r0,q1,r1,c0,c1=map(lambda x:Decimal(str(x)),row[1:]);require(q0*r0==c0 and q1*r1==c1,'fixture multiplication mismatch');base+=c0;current+=c1
require(base==200 and current==244 and (current-base)/base==Decimal('0.22'),'fixture aggregate mismatch')
require((Decimal(1200)-1000)*Decimal('.10')+Decimal(1200)*(Decimal('.12')-Decimal('.10'))==44,'fixture decomposition mismatch')

(ROOT/'VALIDATION.md').write_text('# Validation in progress\n\nThe checker has started; this record is replaced with its actual result on completion.\n', encoding='utf-8')

htmls=list(ROOT.glob('*.html'))+list(ROOT.glob('*/*.html'))
soups={p:BeautifulSoup(p.read_text(), 'html.parser') for p in htmls};link_count=0
for path,soup in soups.items():
    ids=[x['id'] for x in soup.select('[id]')];require(len(ids)==len(set(ids)),str(path.relative_to(ROOT))+' duplicate HTML IDs')
    for a in soup.select('a[href]'):
        href=a['href'];parts=urlsplit(href)
        if parts.scheme or parts.netloc:continue
        link_count+=1;target=(path.parent/unquote(parts.path)).resolve() if parts.path else path
        require(target.is_file(),str(path.relative_to(ROOT))+' broken local link '+href)
        if parts.fragment and target.is_file() and target.suffix=='.html':
            target_soup=soups.get(target) or BeautifulSoup(target.read_text(),'html.parser');require(target_soup.find(id=unquote(parts.fragment)) is not None,str(path.relative_to(ROOT))+' broken anchor '+href)

ppt_checks=[]
for path in sorted((ROOT/'presentations').glob('*.pptx')):
    prs=Presentation(path);bounds=[]
    for n,slide in enumerate(prs.slides,1):
        for sh in slide.shapes:
            if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+50 or sh.top+sh.height>prs.slide_height+50:bounds.append(n)
        require(bool(slide.notes_slide.notes_text_frame.text.strip()),path.name+f' missing notes {n}')
    require(not bounds,path.name+' shapes outside slide '+str(bounds));ppt_checks.append({'file':path.name,'slides':len(prs.slides),'shape_bounds_errors':bounds,'notes_present':True})

browser_checks=[]
with sync_playwright() as p:
    browser=p.chromium.launch(headless=True,args=['--no-sandbox'],executable_path=os.environ.get('CHROMIUM_EXECUTABLE'))
    for key in KEYS:
        page=browser.new_page(viewport={'width':1440,'height':1000});js_errors=[];page.on('pageerror',lambda e:js_errors.append(str(e)))
        navigate(page,ROOT/key/REPORT[key]);page.wait_for_timeout(100)
        total=len(docs[key]['functions']);require(page.locator('.function:visible').count()==total,key+' function visibility')
        page.locator('#function-search').fill('zzzz-no-matching-function');require(page.locator('.function:visible').count()==0,key+' empty search')
        page.locator('#reset-filter').click();require(page.locator('.function:visible').count()==total,key+' reset search')
        group=docs[key]['groups'][0]['id'];page.locator('#function-group').select_option(group);require(page.locator('.function:visible').count()==sum(f['group']==group for f in docs[key]['functions']),key+' group filter')
        page.locator('#reset-filter').click();page.locator('#function-availability').select_option('Preview');require(page.locator('.function:visible').count()==sum(f['availability']=='Preview' for f in docs[key]['functions']),key+' preview filter')
        page.locator('#reset-filter').click();page.locator('#expand-all').click();require(page.locator('main details:not([open])').count()==0,key+' expand disclosures')
        page.locator('#collapse-all').click();before=page.locator('details[open]').count();page.evaluate("window.dispatchEvent(new Event('beforeprint'))");require(page.locator('main details:not([open])').count()==0,key+' print expansion');page.evaluate("window.dispatchEvent(new Event('afterprint'))");require(page.locator('details[open]').count()==before,key+' print restoration')
        screens=ROOT/'qa/screens';screens.mkdir(parents=True,exist_ok=True);page.locator('.function').first.scroll_into_view_if_needed();page.screenshot(path=str(screens/(key+'-desktop.png')),full_page=False)
        page.set_viewport_size({'width':390,'height':844});page.wait_for_timeout(80)
        require(page.evaluate('document.documentElement.scrollWidth<=innerWidth+1'),key+' mobile document overflow')
        page.locator('.function').first.scroll_into_view_if_needed();page.screenshot(path=str(screens/(key+'-mobile.png')),full_page=False)
        require(not js_errors,key+' report JavaScript errors '+str(js_errors));page.close()
        page=browser.new_page(viewport={'width':1440,'height':1000});page.on('pageerror',lambda e:js_errors.append(str(e)));navigate(page,ROOT/key/'presentation.html')
        require(page.locator('.slide:visible').count()==1,key+' deck initial state');page.locator('#next').click();require(page.locator('#counter').inner_text().startswith('2 /'),key+' next slide');page.keyboard.press('ArrowLeft');require(page.locator('#counter').inner_text().startswith('1 /'),key+' keyboard navigation');page.locator('#notes').click();require(page.locator('.slide:visible .speaker-notes').is_visible(),key+' notes toggle');page.locator('#timer-toggle').click();page.wait_for_timeout(1150);require(page.locator('#clock').inner_text()!='00:00',key+' timer');page.locator('#timer-toggle').click();page.locator('#timer-reset').click();require(page.locator('#clock').inner_text()=='00:00',key+' timer reset');page.locator('#all-slides').click();require(page.locator('.slide:visible').count()==len(teaching[key]['slides']),key+' all slides')
        require(not js_errors,key+' deck JavaScript errors '+str(js_errors));page.close();browser_checks.append({'product':key,'report_filters':True,'disclosures_and_print_restoration':True,'mobile_page_overflow':False,'deck_navigation_notes_timer':True,'javascript_errors':js_errors})
    for name in ('index.html','comparison.html'):
        page=browser.new_page(viewport={'width':390,'height':844});navigate(page,ROOT/name);require(page.evaluate('document.documentElement.scrollWidth<=innerWidth+1'),name+' mobile overflow');page.screenshot(path=str(ROOT/'qa/screens'/(name.replace('.html','')+'-mobile.png')));page.close()
    browser.close()

result={'browser_loading':os.environ.get('HTML_LOAD_MODE','file'),'review_date':'2026-10-21','reference_checks':checks,'html_files':len(htmls),'local_links_checked':link_count,'browser_checks':browser_checks,'presentation_checks':ppt_checks,'synthetic_arithmetic':'passed' if not any('fixture' in e for e in errors) else 'failed','errors':errors,'boundary':'Structural, arithmetic and browser checks only. Presentation shape bounds do not prove absence of text overflow; visual rendering is recorded separately. No live AWS execution.'}
(ROOT/'qa').mkdir(exist_ok=True);(ROOT/'qa/validation.json').write_text(json.dumps(result,indent=2)+'\n')
status='passed' if not errors else 'FAILED'
(ROOT/'VALIDATION.md').write_text(f'''# Validation record

Source review date: 21 October 2026. Automated checks: **{status}**.

## What was checked

The checker validates unique function identifiers, group membership, required description fields, separation of availability from activation, official-source hosts, reference resolution, procedure-to-function links and demonstration timing. It verifies the synthetic calculation independently. These checks do not prove that prose is factually correct; source review remains a separate editorial activity.

Generated local links and anchors are checked. Browser tests exercise report search, group and preview filters, disclosures, print-state restoration, mobile page width, slide navigation, notes and elapsed timers. Wide data tables can scroll horizontally within their containers; the page itself must fit the mobile viewport.

PowerPoint files are checked for shape bounds and populated speaker notes. Bounds checks do not detect every text overflow. Rendering and visual inspection, when performed, are recorded separately in `qa/visual-review.json`; absence of that record means no rendered-review claim is made.

See [qa/validation.json](qa/validation.json) for machine-readable results. No live AWS workflow, billing query, scan, resource mutation or memory-recall experiment was executed. No operational improvement is measured.

## Documentation limits

Each product report contains a source register and explicit documentation qualifications. Feature-specific current pages and dated change history take precedence over stale general summaries. The references are not a specification of hidden implementation details or universal assurances about account-specific behavior.

## Publication

The workflow checks before committing each changed file or deletion. It does not make manual historical copies or distribution archives. Previous revisions remain in Git history. Publication to the repository is not deployment to GitHub Pages.
''',encoding='utf-8')
manifest={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(ROOT.rglob('*')) if p.is_file() and '.git' not in p.parts and '__pycache__' not in p.parts and 'screens' not in p.parts and 'rendered' not in p.parts and p.name!='MANIFEST.json'}
(ROOT/'MANIFEST.json').write_text(json.dumps({'algorithm':'sha256','files':manifest},indent=2)+'\n')
print(json.dumps(result,indent=2))
if errors:raise SystemExit(1)
