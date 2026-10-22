#!/usr/bin/env python3
"""Build one canonical guide and a curated presentation per subject.
No historical copies or compressed distributions are created.
"""
from __future__ import annotations
import argparse, hashlib, html, json, re, subprocess
from pathlib import Path
from bs4 import BeautifulSoup, NavigableString
from markdown_it import MarkdownIt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import ImageFont
from lesson_plan import DECKS
ROOT=Path(__file__).resolve().parents[1]
BASE='2b288e7643ccdaa5159a1b6333261b23e971c2ef'
REV='2026-10-22-canonical-guides'
SECTION=re.compile(r'^([DSFA]-(?:M|U)?\d+)\s+[—-]\s+(.*)$')
REF=re.compile(r'^\[([DSFA]S\d+)\]:\s+(https://\S+)',re.M)
MD=MarkdownIt('commonmark',{'html':False}).enable('table')
PALETTE=dict(ink='172C43',navy='142C43',teal='096B68',muted='526579',line='D7E1E8',pale='EFF4F7',white='FFFFFF')
FONT='Liberation Sans'
FONTFILE='/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'
CSS='''
:root{--ink:#172c43;--navy:#142c43;--accent:#096b68;--muted:#526579;--line:#d7e1e8;--pale:#eff4f7}*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:90px}body{margin:0;color:var(--ink);background:#f6f8fa;font:16px/1.65 system-ui,-apple-system,Segoe UI,sans-serif}a{color:#075f98;text-underline-offset:3px}button,input,select{font:inherit;color:inherit;padding:9px 12px;border:1px solid #96a9b8;border-radius:5px;background:white}button{cursor:pointer}button:hover{background:#e9f4f2}:focus-visible{outline:3px solid #b46e08;outline-offset:3px}.topbar{background:var(--navy);color:white;padding:18px 30px;display:flex;align-items:center;justify-content:space-between;gap:18px;flex-wrap:wrap}.topbar a{color:#dceaff}.brand{font-size:13px;letter-spacing:.09em;text-transform:uppercase;font-weight:700}.topbar nav{display:flex;gap:18px;flex-wrap:wrap}.shell{display:grid;grid-template-columns:265px minmax(0,1fr);gap:32px;max-width:1440px;margin:auto;padding:26px 28px}aside{position:sticky;top:20px;max-height:94vh;overflow:auto;align-self:start;padding:4px 14px 12px 0}aside h2{font-size:13px;text-transform:uppercase;letter-spacing:.1em}aside a{display:block;font-size:13px;line-height:1.4;padding:7px 9px;border-left:2px solid var(--line);text-decoration:none;margin:3px 0}aside a:hover{border-color:var(--accent);background:#e7efef}.content{min-width:0;max-width:1000px}.lead{background:white;padding:30px;border:1px solid var(--line);border-top:5px solid var(--accent);border-radius:8px}h1{font-size:clamp(28px,3.4vw,42px);line-height:1.15;letter-spacing:-.025em;margin:0 0 22px}h2{font-size:26px;line-height:1.25;margin:0 0 20px}h3{font-size:20px;line-height:1.35;margin:26px 0 12px}p{margin:0 0 17px}.report-section{background:white;border:1px solid var(--line);padding:30px;margin:25px 0;border-radius:8px;min-width:0}.section-label{font-size:11px;letter-spacing:.1em;font-weight:700;color:var(--accent);text-transform:uppercase;margin-bottom:9px}.controls{display:flex;gap:10px;flex-wrap:wrap;background:#f6f8fa;padding:15px 0;position:sticky;top:0;z-index:3}.controls input{flex:2;min-width:180px}.controls select{flex:1;min-width:130px}.status{font-size:13px;color:var(--muted);margin:0 0 10px}.table-wrap{overflow-x:auto;margin:20px 0}table{border-collapse:collapse;width:100%;font-size:14px;line-height:1.55}th,td{padding:13px 15px;border:1px solid var(--line);text-align:left;vertical-align:top;min-width:90px}th{background:#e9f0f4;font-weight:700}tr:nth-child(even) td{background:#f8fafb}pre{background:#11283d;color:#eaf2f8;padding:20px;overflow:auto;border-radius:6px;font-size:13px;line-height:1.6}code{font-size:.9em;overflow-wrap:anywhere}blockquote{border-left:4px solid var(--accent);margin:20px 0;padding:5px 20px;background:#eef6f4}li{margin:7px 0}.ref-link{font-size:12px;font-weight:600}footer{font-size:13px;color:var(--muted);padding:22px 0}.toc-toggle{display:none}[hidden]{display:none!important}.cards{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:24px}.card{background:white;border:1px solid var(--line);padding:26px;border-top:4px solid var(--accent);border-radius:8px}.card h2{font-size:24px}.card a{margin-right:12px}.guide-index{max-width:1180px;margin:35px auto;padding:0 24px}.small{font-size:13px;color:var(--muted)}.deck{display:none;max-width:1320px;margin:20px auto;padding:0 24px}.presenting .shell{display:none}.presenting .deck{display:block}.deck-controls{display:flex;align-items:center;gap:10px;flex-wrap:wrap;margin:15px 0}.deck-controls select{min-width:200px;flex:1}.slide{background:white;border:1px solid var(--line);border-top:6px solid var(--accent);padding:25px 34px;min-height:640px;display:flex;flex-direction:column}.slide .kicker{font-size:12px;letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}.slide h2{font-size:clamp(26px,3vw,39px);margin:12px 0 28px;max-width:1100px}.slide-body{flex:1;display:flex;align-items:center;min-height:310px}.slide table{font-size:19px}.slide th,.slide td{padding:15px;line-height:1.4}.flow{display:flex;gap:15px;width:100%;align-items:stretch}.node{flex:1;min-width:0;background:var(--pale);border-top:4px solid var(--accent);padding:20px;position:relative}.node:not(:last-child):after{content:'→';position:absolute;right:-15px;top:43%;font-size:23px;color:var(--accent)}.node strong{font-size:21px;display:block;margin-bottom:15px;line-height:1.25}.node p{font-size:19px;line-height:1.5}.columns{display:flex;gap:22px;width:100%}.column{flex:1;border-top:4px solid var(--accent);padding:22px;background:var(--pale);min-width:0}.column strong{font-size:23px;display:block;margin-bottom:18px}.column p{font-size:20px;line-height:1.5}.takeaway{background:var(--navy);color:white;padding:15px 20px;font-size:18px;line-height:1.4;margin:25px 0 10px}.slide-foot{font-size:12px;color:var(--muted);display:flex;justify-content:space-between;gap:15px;flex-wrap:wrap}.notes{background:#f0f4f6;border:1px solid var(--line);padding:24px;margin-top:16px;max-height:500px;overflow:auto}.notes h2{font-size:22px}.notes h3{font-size:17px}.notes table{font-size:13px}.presenting .topbar{position:static}.timer{font-variant-numeric:tabular-nums;font-size:16px;font-weight:700}.source-foot{overflow-wrap:anywhere}
@media(max-width:900px){.shell{grid-template-columns:1fr;padding:16px;gap:0}aside{position:static;max-height:250px;background:white;border:1px solid var(--line);padding:10px;margin-bottom:16px}.toc-toggle{display:inline-block}.report-section,.lead{padding:22px}.controls{position:static}.cards{grid-template-columns:1fr}.slide{padding:20px;min-height:0}.flow,.columns{flex-direction:column}.node:not(:last-child):after{display:none}.slide table{font-size:16px}.slide-body{min-height:0}.node strong,.column strong{font-size:20px}.node p,.column p{font-size:18px}.deck{padding:0 12px}.slide h2{font-size:27px}.slide .table-wrap{max-width:100%}.topbar{padding:16px}.takeaway{font-size:16px}}
@media print{body{background:white;font-size:10pt}.topbar,aside,.controls,.deck-controls,.status,.notes,footer{display:none!important}.shell{display:block;padding:0;max-width:none}.content{max-width:none}.lead,.report-section{border:0;padding:0;margin:15px 0;break-inside:auto}.report-section[hidden]{display:block!important}.table-wrap{overflow:visible}pre{white-space:pre-wrap;color:black;background:#eee}table{font-size:9pt}.deck{padding:0;max-width:none}.presenting .slide{display:flex!important;break-after:page;min-height:0;height:95vh;border:0}.presenting .slide[hidden]{display:flex!important}.slide h2{font-size:27pt}.slide table{font-size:16pt}.slide .flow,.slide .columns{flex-direction:row}.slide .node p,.slide .column p{font-size:15pt}.slide .node strong,.slide .column strong{font-size:17pt}.takeaway{color:black;background:#eee}.source-foot{font-size:9pt}}
'''
JS='''(()=>{const sections=[...document.querySelectorAll('.report-section')],q=document.querySelector('#search'),kind=document.querySelector('#kind'),status=document.querySelector('#status');
function filter(){let n=0;sections.forEach(s=>{let ok=(!q.value||s.textContent.toLowerCase().includes(q.value.toLowerCase()))&&(!kind.value||kind.value===s.dataset.kind);s.hidden=!ok;n+=ok;});status.textContent=n+' of '+sections.length+' sections shown';}
if(q){q.addEventListener('input',filter);kind.addEventListener('change',filter);document.querySelector('#reset').onclick=()=>{q.value='';kind.value='';filter();};filter();}
function reveal(){const s=document.getElementById(decodeURIComponent(location.hash.slice(1)));if(s&&s.classList.contains('report-section')){if(q){q.value='';kind.value='';filter();}document.body.classList.remove('presenting');s.scrollIntoView();}}
window.addEventListener('hashchange',reveal);reveal();
const slides=[...document.querySelectorAll('.slide')],select=document.querySelector('#slide-select');let current=0,start=0,elapsed=0,running=false;
function clock(v){return Math.floor(v/60)+':'+String(v%60).padStart(2,'0');}
function show(i){if(!slides.length)return;current=Math.max(0,Math.min(slides.length-1,i));slides.forEach((s,j)=>s.hidden=j!==current);select.value=current;document.querySelectorAll('.notes').forEach((n,j)=>n.hidden=j!==current||!document.querySelector('#notes-toggle').checked);}
if(slides.length){document.querySelector('#present').onclick=()=>{document.body.classList.add('presenting');show(current);window.scrollTo(0,0);};document.querySelector('#read').onclick=()=>document.body.classList.remove('presenting');document.querySelector('#prev').onclick=()=>show(current-1);document.querySelector('#next').onclick=()=>show(current+1);select.onchange=()=>show(Number(select.value));document.querySelector('#notes-toggle').onchange=()=>show(current);document.querySelector('#timer-toggle').onclick=()=>{if(running){elapsed+=Math.floor((Date.now()-start)/1000);running=false;}else{start=Date.now();running=true;}document.querySelector('#timer-toggle').textContent=running?'Pause timer':'Start timer';};document.querySelector('#timer-reset').onclick=()=>{elapsed=0;start=Date.now();};setInterval(()=>{document.querySelector('#timer').textContent=clock(elapsed+(running?Math.floor((Date.now()-start)/1000):0));},250);show(0);}
document.querySelectorAll('[data-print]').forEach(b=>b.onclick=()=>window.print());document.addEventListener('keydown',e=>{if(['INPUT','SELECT','TEXTAREA'].includes(document.activeElement.tagName))return;if(e.key==='/'&&q){e.preventDefault();q.focus();}if(document.body.classList.contains('presenting')){if(e.key==='ArrowRight'){e.preventDefault();show(current+1);}if(e.key==='ArrowLeft'){e.preventDefault();show(current-1);}if(e.key==='Home'){e.preventDefault();show(0);}if(e.key==='End'){e.preventDefault();show(slides.length-1);}if(e.key==='Escape')document.body.classList.remove('presenting');}});
})();'''

def write(rel,text):
    p=ROOT/rel;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(text,encoding='utf-8')
def h(s):return html.escape(str(s))
def git(*args):return subprocess.check_output(['git','-C',str(ROOT),*args],text=True).strip()
def fixes():
    p=ROOT/'guides/devops.md'
    if not p.exists():return
    t=p.read_text()
    t=t.replace('## D-M02 — Corrective feedback and custom-store access','## D-M02 — Corrective feedback produces a separate lesson')
    if '## D-M07' not in t:t=t.replace('A new custom agent has no store attachments.','## D-M07 — Memory attachment and lifecycle\n\nA new custom agent has no store attachments.',1)
    t=t.replace(' description: placeholder','description: Investigate checkout timeouts involving the inventory dependency.')
    old='Do not use that placeholder as a real skill. A useful description would state: “Use when checkout requests time out while waiting for the inventory dependency; compare upstream response time, connection errors and deployment timestamps.” The procedure should then identify the actual connected tools and expected evidence. This contrast illustrates the matching problem without inventing a native tool name. A procedure cannot confer missing tool permissions.'
    new='Add a procedure that compares upstream response time, connection errors and deployment timestamps using the actually connected tools. Record failed reads and distinguish observations from hypotheses. The example frontmatter illustrates applicability; it is not a complete environment-specific runbook. A procedure cannot confer missing tool permissions.'
    t=t.replace(old,new)
    def cite(m):
        text=m.group(1);ids=list(dict.fromkeys(re.findall(r'DS\d+',text)))
        return ' '.join('['+i+']' for i in ids)
    t=re.sub(r'\[(DS\d+[^\]\n]*[:;][^\]\n]*)\]',cite,t)
    p.write_text(t,encoding='utf-8')

def load_guides():
    guides={}
    for k in DECKS:
        raw=(ROOT/f'guides/{k}.md').read_text();refs=dict(REF.findall(raw));soup=BeautifulSoup(MD.render(raw),'html.parser')
        sections={};intro=[];current=None
        for el in list(soup.contents):
            if getattr(el,'name',None)=='h2':
                m=SECTION.match(el.get_text())
                if not m:raise ValueError('Unidentified section '+el.get_text())
                current=m[1];sections[current]={'title':m[2],'nodes':[el]}
            elif current:sections[current]['nodes'].append(el)
            else:intro.append(el)
        for sid,s in sections.items():
            body=''.join(str(x) for x in s['nodes']);s['html']=body;s['text']=BeautifulSoup(body,'html.parser').get_text(' ',strip=True)
            s['source_ids']=sorted(set(re.findall(r'\b[DSFA]S\d+\b',s['text'])) & set(refs))
        guides[k]={'raw':raw,'refs':refs,'sections':sections,'intro':''.join(map(str,intro))}
    return guides

def kind(sid,title):
    if '-M' in sid:return 'mechanism'
    if '-U' in sid:return 'example'
    if any(x in title.lower() for x in ('coverage','evidence notes','source-review','unresolved')):return 'reference'
    if any(x in title.lower() for x in ('limits','diagnostic','troubleshooting')):return 'troubleshooting'
    if any(x in title.lower() for x in ('model','resources before')):return 'model'
    return 'procedure'

def clean_html(text,ids):
    soup=BeautifulSoup(text,'html.parser')
    for table in soup.select('table'):
        wrap=soup.new_tag('div',attrs={'class':'table-wrap'});table.wrap(wrap)
    for a in soup.select('a[href]'):
        if a.get('href','').startswith('https://docs.aws.amazon.com'):a['class']='ref-link';a['rel']='noopener'
    for node in list(soup.find_all(string=True)):
        if node.find_parent(['a','pre','code','h1','h2','h3','script','style']):continue
        parts=re.split(r'\b([DSFA]-(?:M|U)?\d+)\b',str(node))
        if len(parts)==1:continue
        for i,x in enumerate(parts):
            if i%2 and x in ids:
                a=soup.new_tag('a',href=ids[x]);a.string=x;node.insert_before(a)
            else:node.insert_before(NavigableString(x))
        node.extract()
    return str(soup)

def frame(title,body):
    return '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+h(title)+'</title><style>'+CSS+'</style></head><body><header class="topbar"><span class="brand">Frontier agents · engineering guides</span><nav><a href="../index.html">Collection</a><a href="../EDITORIAL_STANDARD.md">Editorial standard</a><a href="../VALIDATION.md">Validation</a></nav></header>'+body+'<script>'+JS+'</script></body></html>'

def slide_body(sl):
    rows=sl['data'];kind_=sl['kind']
    if kind_ in ('table','demo'):
        if kind_=='demo':rows=[['Time','Action','Prepared evidence'],*rows]
        return '<div class="table-wrap"><table><thead><tr>'+''.join('<th>'+h(x)+'</th>' for x in rows[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+h(x)+'</td>' for x in row)+'</tr>' for row in rows[1:])+'</tbody></table></div>'
    cls='flow' if kind_=='flow' else 'columns';item='node' if kind_=='flow' else 'column'
    return '<div class="'+cls+'">'+''.join('<div class="'+item+'"><strong>'+h(a)+'</strong><p>'+h(b)+'</p></div>' for a,b in rows)+'</div>'

def schedule(d):
    elapsed=0;out=[]
    for s in d['slides']:
        out.append((elapsed,elapsed+s['seconds']));elapsed+=s['seconds']
    return out

def stamp(v):return f'{v//60:02d}:{v%60:02d}'

def make_html(k,d,guides):
    g=guides[k];ids={sid:(('#'+sid) if key==k else key+'.html#'+sid) for key,v in guides.items() for sid in v['sections']}
    toc=''.join('<a href="#'+sid+'">'+h(sid+' · '+s['title'])+'</a>' for sid,s in g['sections'].items())
    report=''
    for sid,s in g['sections'].items():
        report+='<section class="report-section" id="'+sid+'" data-kind="'+kind(sid,s['title'])+'"><div class="section-label">'+h(kind(sid,s['title']))+'</div>'+clean_html(s['html'],ids)+'</section>'
    controls='<div class="controls"><input id="search" aria-label="Search guide" placeholder="Search the guide · / to focus"><select id="kind" aria-label="Section type"><option value="">All sections</option>'+''.join('<option value="'+x+'">'+x.title()+'</option>' for x in ('model','procedure','mechanism','troubleshooting','example','reference'))+'</select><button id="reset">Reset</button><button id="present">Presentation</button><button data-print>Print</button></div><p id="status" class="status" aria-live="polite"></p>'
    slides='';notes='';transcript=['# '+d['title']+' — presenter reference', '', 'The schedule includes prepared-result exercises. It is not a promise of live agent completion time. Questions and breaks are additional. Canonical guide sections are reproduced in speaker notes as reference, not as a required verbatim script.', '']
    for i,(s,times) in enumerate(zip(d['slides'],schedule(d)),1):
        links=' · '.join('<a href="#'+r+'">'+r+'</a>' for r in s['refs'])
        foot=f'{i} / {len(d["slides"])} · {stamp(times[0])}–{stamp(times[1])} · {s["seconds"]//60} min'
        slides+='<article class="slide" data-seconds="'+str(s['seconds'])+'" data-kind="'+s['kind']+'"'+(' hidden' if i>1 else '')+'><div class="kicker">'+h(d['title']+' · '+s['kind'])+'</div><h2>'+h(s['title'])+'</h2><div class="slide-body">'+slide_body(s)+'</div><div class="takeaway">'+h(s['message'])+'</div><div class="slide-foot"><span>'+foot+'</span><span>'+links+'</span></div></article>'
        refhtml=''.join(g['sections'][r]['html'] for r in s['refs'])
        notes+='<div class="notes" hidden>'+clean_html(refhtml,ids)+'</div>'
        transcript+=['## '+str(i)+'. '+s['title'],f'**Schedule:** {stamp(times[0])}–{stamp(times[1])}. **Type:** {s["kind"]}.',s['message'],'**Display content:**',*[' | '.join(x) for x in s['data']],'**Canonical reference:** '+', '.join('['+r+'](../guides/'+k+'.html#'+r+')' for r in s['refs']),'']
    total=sum(s['seconds'] for s in d['slides'])//60;demo=sum(s['seconds'] for s in d['slides'] if s['kind']=='demo')//60
    deck='<div class="deck"><div class="deck-controls"><button id="read">Read guide</button><button id="prev">Previous</button><select id="slide-select" aria-label="Select slide">'+''.join('<option value="'+str(i)+'">'+str(i+1)+'. '+h(s['title'])+'</option>' for i,s in enumerate(d['slides']))+'</select><button id="next">Next</button><label><input id="notes-toggle" type="checkbox"> Notes</label><button id="timer-toggle">Start timer</button><button id="timer-reset">Reset timer</button><span id="timer" class="timer">0:00</span><button data-print>Print slides</button></div><p class="small">'+str(total)+' minutes including '+str(demo)+' minutes of prepared-result exercise. Arrow keys navigate; Escape returns to guide.</p>'+slides+notes+'</div>'
    body='<div class="shell"><aside><h2>Contents</h2>'+toc+'</aside><main class="content"><div class="lead">'+clean_html(g['intro'],ids)+'<p class="small"><a href="'+k+'.md">Markdown source</a> · <a href="../presentations/'+k+'.pptx">Editable PowerPoint</a> · <a href="../presentations/'+k+'-notes.md">Presenter reference</a></p></div>'+controls+report+'<footer>Documentation review is separate from configuration and runtime testing. No live AWS execution was performed.</footer></main></div>'+deck
    write(f'guides/{k}.html',frame(d['title'],body));write(f'presentations/{k}-notes.md','\n\n'.join(transcript)+'\n')

def fitted(text,w,hgt,size,minimum):
    try:
        for s in range(int(size),int(minimum)-1,-1):
            font=ImageFont.truetype(FONTFILE,s*4);limit=(w-.08)*72*4;lines=0
            for para in str(text).split('\n'):
                acc='';lines+=1
                for word in para.split():
                    candidate=(acc+' '+word).strip()
                    if font.getlength(candidate)>limit and acc:lines+=1;acc=word
                    else:acc=candidate
            if lines*s*1.18 <= (hgt-.02)*72:return s
        raise ValueError('Text requires a split: '+str(text)[:100])
    except OSError:return minimum

def box(sl,x,y,w,hh,fill,line=None):
    sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(hh));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background() if not line else None
    if line:sh.line.color.rgb=RGBColor.from_string(line)
    return sh

def txt(sl,text,x,y,w,hh,size=20,bold=False,color='172C43',minimum=16):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(hh));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    fs=fitted(text,w,hh,size,minimum)
    for j,line in enumerate(str(text).split('\n')):
        p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(fs);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0);p.line_spacing=1.12
    return sh

def make_pptx(k,d,g):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    prs.core_properties.title=d['title']+' — engineering presentation';prs.core_properties.subject='Canonical guide '+REV;prs.core_properties.author='Engineering documentation'
    for i,(s,times) in enumerate(zip(d['slides'],schedule(d)),1):
        sl=prs.slides.add_slide(prs.slide_layouts[6]);box(sl,0,0,13.333,.09,PALETTE['teal']);txt(sl,d['title'].upper(),.55,.25,12,.24,11,True,PALETTE['muted'],11)
        txt(sl,s['title'],.55,.72,12.15,.95,30,True,minimum=26)
        if s['kind'] in ('table','demo'):
            rows=s['data'] if s['kind']=='table' else [['Time','Action','Prepared evidence'],*s['data']]
            n=len(rows);cols=len(rows[0]);rowheight=4.15/n
            widths=([1.2,3.1,7.95] if s['kind']=='demo' else ([3.0,9.25] if cols==2 else ([2.6,4.3,5.35] if cols==3 else [2.0,3.4,3.4,3.45])))
            for r,row in enumerate(rows):
                xx=.55
                for c,val in enumerate(row):
                    fill=PALETTE['navy'] if r==0 else (PALETTE['pale'] if r%2 else 'FFFFFF');box(sl,xx,1.92+r*rowheight,widths[c],rowheight,fill,PALETTE['line']);txt(sl,val,xx+.14,2.04+r*rowheight,widths[c]-.28,rowheight-.21,17 if r==0 else 19,r==0,'FFFFFF' if r==0 else PALETTE['ink'],minimum=15);xx+=widths[c]
        else:
            n=len(s['data']);gap=.28;w=(12.25-gap*(n-1))/n
            for j,(label,body) in enumerate(s['data']):
                x=.55+j*(w+gap);box(sl,x,2.0,w,3.95,PALETTE['pale']);box(sl,x,2,w,.05,PALETTE['teal']);txt(sl,label,x+.18,2.25,w-.36,.85,22,True,minimum=18);txt(sl,body,x+.18,3.35,w-.36,2.2,21,minimum=17)
                if s['kind']=='flow' and j<n-1:txt(sl,'→',x+w+.015,3.65,gap-.03,.4,19,True,PALETTE['teal'],17)
        box(sl,.55,6.24,12.25,.63,PALETTE['navy']);txt(sl,s['message'],.75,6.38,11.85,.37,17,False,'FFFFFF',15)
        txt(sl,' · '.join(s['refs']),.55,7.05,7.0,.2,10,False,PALETTE['muted'],10);txt(sl,f'{i:02d} / {len(d["slides"]):02d}     {stamp(times[0])}–{stamp(times[1])}  |  {s["seconds"]//60} min',8.0,7.05,4.8,.2,10,False,PALETTE['muted'],10)
        notes=[f'SLIDE {i}. {s["title"]}',f'Budget {s["seconds"]//60} minutes; cumulative {stamp(times[0])}–{stamp(times[1])}.',s['message'],'Prepared-result exercise; no live result claimed.' if s['kind']=='demo' else 'Use the display as the speaking outline. The following canonical guide text is reference, not a mandatory verbatim script.']
        for ref in s['refs']:
            sec=g['sections'][ref];notes.extend([ref+' — '+sec['title'],sec['text']]);notes.extend(sid+' '+g['refs'][sid] for sid in sec['source_ids'])
        sl.notes_slide.notes_text_frame.text='\n\n'.join(notes)
    sl=prs.slides.add_slide(prs.slide_layouts[6]);box(sl,0,0,13.333,.09,PALETTE['teal']);txt(sl,'Evidence and use of this presentation',.55,.7,12,1,32,True,minimum=28)
    for i,(label,body) in enumerate([('Canonical guide',f'guides/{k}.md — resource model, procedures, mechanisms, troubleshooting and worked example.'),('Source traceability','Every content slide names guide sections. Speaker notes include the related guide text and official AWS source URLs.'),('Delivery boundary','Presentation time includes the prepared-result exercise. No live AWS execution, measured benefit or exhaustive product coverage is asserted.'),('Scope decisions','Read the coverage and source-version notes before extending this material to another feature, Region or provider.')]):
        txt(sl,label,.6,2.0+i*.98,2.3,.65,18,True,minimum=16);txt(sl,body,3.15,2.0+i*.98,9.45,.72,20,minimum=17)
    txt(sl,'UNTIMED APPENDIX · '+REV,.55,7.05,12,.22,11,False,PALETTE['muted'],11)
    sl.notes_slide.notes_text_frame.text='Evidence appendix, outside the timed programme.\n\n'+'\n'.join(i+' '+u for i,u in g['refs'].items())
    p=ROOT/f'presentations/{k}.pptx';p.parent.mkdir(exist_ok=True);prs.save(p)

def cleanup():
    if not (ROOT/'.git').exists():return []
    owned=('devops/','security/','finops/','architecture/','platforms/','mechanisms/','content/','presentations/','qa/','demo-assets/','assets/')
    oldtools={'tools/assemble_content.py','tools/build_mechanisms.py','tools/build_presentations.py','tools/build_suite.py','tools/editorial_lint.py','tools/qa_format.py','tools/qa_mechanisms_browser.py','tools/validate_scope.py','tools/build_platform_references.py','tools/finalize_platform_publication.py'}
    removed=[]
    for rel in git('ls-tree','-r','--name-only',BASE).splitlines():
        if rel.startswith(owned) or rel in oldtools or rel in {'FORMAT_IMPLEMENTATION.md','style-guide.html'}:
            p=ROOT/rel
            if p.is_file():p.unlink();removed.append(rel)
    for rel in sorted(owned,key=len,reverse=True):
        p=ROOT/rel
        if p.exists():
            for sub in sorted(p.rglob('*'),key=lambda x:len(x.parts),reverse=True):
                if sub.is_dir() and not any(sub.iterdir()):sub.rmdir()
            if not any(p.iterdir()):p.rmdir()
    return removed

def root_pages(guides):
    cards='';rows=[];total=demo=0
    for k,d in DECKS.items():
        mins=sum(s['seconds'] for s in d['slides'])//60;dm=sum(s['seconds'] for s in d['slides'] if s['kind']=='demo')//60;total+=mins;demo+=dm
        cards+='<article class="card"><h2>'+h(d['title'])+'</h2><p>'+h(d['subtitle'])+'</p><p class="small">'+str(mins)+' minutes including '+str(dm)+' minutes of prepared-result exercise</p><a href="guides/'+k+'.html">Read / present</a><a href="guides/'+k+'.md">Markdown</a><a href="presentations/'+k+'.pptx">PowerPoint</a></article>'
        rows.append(f'| {d["title"]} | [Guide](guides/{k}.html) / [Markdown](guides/{k}.md) | [Slides](presentations/{k}.pptx) / [Notes](presentations/{k}-notes.md) | {mins} min including {dm} min exercise |')
    index=frame('AWS agent engineering guides','<main class="guide-index"><h1>Configure the platform. Understand the mechanism. Verify the result.</h1><p>One canonical guide per platform, plus a comparison of the three control systems. Product facts, configuration instructions and synthetic exercises are explicitly distinguished.</p><div class="cards">'+cards+'</div><h2>Evidence and scope</h2><p>Reviewed 22 October 2026. Read the <a href="guides/coverage.md">coverage map</a>, <a href="VALIDATION.md">validation boundary</a> and <a href="EDITORIAL_STANDARD.md">editorial standard</a>. The '+str(total)+'-minute programme includes '+str(demo)+' exercise minutes; questions and breaks are additional.</p><p>The separate <a href="applications/aws-validation-console/README.md">operational-readiness console application</a> is preserved and is not evidence that these AWS workflows were executed.</p><p>No server is needed to read the HTML locally. Repository publication is not GitHub Pages deployment. Git history retains earlier revisions; no manual archives are created.</p></main>')
    index=index.replace('href="../index.html"','href="index.html"').replace('href="../EDITORIAL_STANDARD.md"','href="EDITORIAL_STANDARD.md"').replace('href="../VALIDATION.md"','href="VALIDATION.md"');write('index.html',index)
    write('README.md','# AWS agent engineering guides\n\nOne canonical source per subject. Clone this repository and open `index.html`; HTML includes searchable reading and presentation modes. GitHub file preview does not execute HTML. No website deployment is implied.\n\n| Subject | Guide | Presentation | Timing |\n|---|---|---|---|\n'+'\n'.join(rows)+f'\n\nTotal: **{total} minutes including {demo} minutes of prepared-result exercises**. Questions and breaks are additional. These are new schedules, not the superseded scenario programme.\n\n## Architecture of the publication\n\n`guides/*.md` is the canonical descriptive source. `tools/lesson_plan.py` is the curated speaking outline; each slide references canonical sections, which supply its speaker-note evidence. HTML and editable PowerPoint are generated by `tools/docs_build.py`. There are no separate competing capability or mechanism catalogs.\n\nRead [coverage](guides/coverage.md), [editorial rules](EDITORIAL_STANDARD.md), [validation](VALIDATION.md), and [publishing](PUBLISHING.md). Older editions are retained only through Git history.\n\n## Rebuild\n\n```sh\npython -m pip install -r tools/docs-requirements.txt\npython -m playwright install chromium\npython tools/docs_build.py\npython tools/docs_check.py --browser\n```\n\nThe build replaces only the prior documentation-owned files listed from a pinned Git baseline; it preserves the [separate readiness-console application](applications/aws-validation-console/README.md) and its workflow. No live AWS calls or external publication tasks are executed by this documentation build.\n')
    write('EDITORIAL_STANDARD.md','# Editorial and technical acceptance standard\n\nThe unit of documentation is the product resource or mechanism, not an arbitrary fixed number of steps.\n\nA capability explains a reusable supported ability. An operation acts on a named object. A mechanism explains documented triggers, state, execution and effects. A use case applies these to a concrete problem. A broad product function may contain several capabilities.\n\nEvery procedure must state meaningful prerequisites, configuration decisions, expected confirmation and failure interpretation. Lifecycle instructions distinguish disable, cancel, soft delete, forget and permanent deletion where the product does. Steps and diagrams have as many elements as the subject requires.\n\nCite the official source near the product claim. Source registers identify reviewed topics and conflicts between retrieved page versions. Missing documented implementation details remain unknown. Synthetic examples identify invented inputs, supply enough evidence to support their conclusion and include a failure branch.\n\nUse direct descriptive prose. Do not replace product explanation with repeated generic cautions or speculative benefits. Present recommendations as engineering guidance, not native features.\n\nAutomated checks establish link, schema, fixture, browser and slide-layout properties only. They cannot certify factual completeness, successful AWS setup, diagnostic correctness, security coverage or savings. Coverage maps state detailed, linked and deferred areas.\n\nOne canonical guide per platform. Globally unique section identifiers. Presentations link to those sections. No manual archive directories or committed distribution packages; Git history is the version record.\n')
    write('SCOPE.md','# Scope\n\nAWS DevOps Agent, AWS Security Agent and AWS FinOps Agent, with a shared architectural comparison. The independently developed readiness-console application is preserved unchanged.\n\nThis is an engineering synthesis of selected official user-guide topics, not a complete API reference, an AWS publication, an exhaustive source audit, or a live service assessment. See each guide’s explicit coverage and source-version decisions.\n')
    write('PUBLISHING.md','# Publication\n\nEdit canonical Markdown and the curated lesson plan. Rebuild and validate before changing `documentation-build.json` to request a publication run. The workflow renders the guides and slides, runs the documented checks, and commits each changed document or deletion separately. It does not create ZIP distributions or snapshot directories.\n\nThe documentation-owned paths are isolated from `applications/` and `.github/workflows/validation-console.yml`. Review the final Git diff and preservation check. No force push or history rewrite is part of publication.\n')
    write('CHANGELOG.md','# Change log\n\n## 22 October 2026 — canonical engineering guides\n\nReplaced overlapping scenario, mechanism and platform-reference editions with one canonical guide per platform and a shared comparison. Added resource models, configuration decisions, lifecycle distinctions, diagnostic tables, source-version decisions and evidence-bearing synthetic exercises. Rebuilt presentations as curated teaching sequences referencing canonical sections. Removed obsolete documentation generators and duplicated editions; Git preserves history. The separate readiness-console application and workflow are unchanged.\n\nThis edition is not a record of live AWS execution. See VALIDATION.md for observed build/test results and limitations.\n\n')
    coverage=['# Coverage and claim-to-source map','','Review date: 22 October 2026. This map enumerates canonical sections and their cited official source identifiers. It is not a percentage of all AWS functionality. Detailed/linked/deferred boundaries and source-version choices are stated in each guide.','']
    for k,g in guides.items():
        coverage+=['## '+DECKS[k]['title'],'','| Canonical section | Subject | Cited source identifiers |','|---|---|---|']
        for sid,s in g['sections'].items():coverage.append('| ['+sid+']('+k+'.html#'+sid+') | '+s['title']+' | '+', '.join(s['source_ids'])+' |')
        coverage+=['']
    write('guides/coverage.md','\n'.join(coverage)+'\n')
    redirects={'devops/validated-report.html':'devops','devops/initial-report.html':'devops','security/report.html':'security','finops/report.html':'finops','architecture/common-patterns.html':'architecture',**{f'platforms/{k}.html':k for k in ('devops','security','finops')},**{f'mechanisms/{k}.html':k for k in ('devops','security','finops')}}
    for rel,k in redirects.items():write(rel,'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Guide moved</title></head><body><h1>Canonical guide</h1><p>This entry point now links to the single maintained guide. It contains no separate edition.</p><a href="../guides/'+k+'.html">Open '+h(DECKS[k]['title'])+'</a></body></html>')
    for rel in ('platforms/index.html','mechanisms/index.html','mechanisms/overview.html'):
        dest='../guides/architecture.html' if 'overview' in rel else '../index.html';write(rel,'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Collection moved</title></head><body><p>Use the canonical collection.</p><a href="'+dest+'">Open current guides</a></body></html>')

def commit():
    raw=subprocess.check_output(['git','status','--porcelain=v1','-z','--untracked-files=all'],cwd=ROOT);paths=[]
    for record in raw.split(b'\0'):
        if not record:continue
        if b'R' in record[:2] or b'C' in record[:2]:raise RuntimeError('Unexpected rename; inspect publication diff')
        path=record[3:].decode()
        if path.startswith('applications/') or path=='.github/workflows/validation-console.yml':raise RuntimeError('Unrelated application change '+path)
        paths.append(path)
    for path in sorted(set(paths)):
        subprocess.run(['git','add','--',path],cwd=ROOT,check=True);subprocess.run(['git','commit','-m','docs: publish '+path],cwd=ROOT,check=True)

def main():
    a=argparse.ArgumentParser();a.add_argument('--commit',action='store_true');args=a.parse_args()
    if args.commit:return commit()
    cleanup();fixes();guides=load_guides()
    for k,d in DECKS.items():
        for s in d['slides']:
            assert all(r in guides[k]['sections'] for r in s['refs']),(k,s['title'])
        make_html(k,d,guides);make_pptx(k,d,guides[k])
    root_pages(guides)
    print(json.dumps({'revision':REV,'guides':list(guides),'timed_slides':sum(len(d['slides']) for d in DECKS.values()),'minutes':sum(s['seconds'] for d in DECKS.values() for s in d['slides'])//60},indent=2))
if __name__=='__main__':main()
