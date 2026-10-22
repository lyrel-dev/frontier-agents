#!/usr/bin/env python3
"""Check publication integrity and executable teaching fixtures, not live AWS behavior."""
from __future__ import annotations
import argparse, hashlib, json, re, shutil, subprocess, tempfile
from decimal import Decimal as D
from pathlib import Path
from urllib.parse import urlsplit, unquote
from zipfile import ZipFile
from bs4 import BeautifulSoup
from pptx import Presentation
import yaml
from docs_build import ROOT,REV,DECKS,load_guides,write,MD,git,SECTION,REF
ROOT_FILES={'README.md','index.html','EDITORIAL_STANDARD.md','SCOPE.md','PUBLISHING.md','CHANGELOG.md','VALIDATION.md','documentation-build.json'}
DOC_DIRS={'guides','presentations','examples','qa','devops','security','finops','architecture','platforms','mechanisms'}
DOC_TOOLS={'tools/docs_build.py','tools/docs_check.py','tools/lesson_plan.py','tools/docs-requirements.txt'}
def publication_files():
    for p in sorted(ROOT.rglob('*')):
        if not p.is_file():continue
        rel=p.relative_to(ROOT)
        if rel.as_posix() in ROOT_FILES|DOC_TOOLS or rel.parts[0] in DOC_DIRS:yield p

def check(browser=False):
    errors=[];g=load_guides();ids=set();sectioncount=0;source_urls=set();links=0;slides=0;browser_results=[];fixture_results=[]
    for k,v in g.items():
        headings=re.findall(r'^## ([DSFA]-(?:M|U)?\d+) ',v['raw'],re.M)
        if len(headings)!=len(set(headings)):errors.append(k+': duplicate section')
        if ids.intersection(headings):errors.append(k+': globally reused ID')
        ids.update(headings);sectioncount+=len(headings)
        body=REF.sub('',v['raw'])
        cited=set(re.findall(r'\[([DSFA]S\d+)\]',body))
        if cited-set(v['refs']):errors.append(k+': undefined citation')
        if re.search(r'\[[DSFA]S\d+[^\]\n]*[:;]',body):errors.append(k+': unresolved extended citation')
        for sid,url in v['refs'].items():
            if urlsplit(url).hostname!='docs.aws.amazon.com':errors.append(sid+': nonofficial source')
            source_urls.add(url)
        for token in MD.parse(v['raw']):
            if token.type=='fence' and token.info.strip()=='yaml':
                try:yaml.safe_load(token.content)
                except yaml.YAMLError:errors.append(k+': invalid YAML example')
            if token.type=='fence' and token.content.startswith('---\nname:'):
                try:yaml.safe_load(token.content.split('---',2)[1])
                except yaml.YAMLError:errors.append(k+': invalid skill frontmatter')
        for s in DECKS[k]['slides']:
            if not set(s['refs'])<=set(v['sections']):errors.append(k+': invalid slide cross-reference')
            if s['seconds']<=0:errors.append(k+': invalid time budget')
            if s['kind']=='demo':
                intervals=[re.match(r'(\d+)–(\d+) min',r[0]) for r in s['data']]
                if not all(intervals):errors.append(k+': invalid exercise interval')
                elif int(intervals[-1][2])*60!=s['seconds'] or any(int(a[2])!=int(b[1]) for a,b in zip(intervals,intervals[1:])):errors.append(k+': exercise time mismatch')
        p=Presentation(ROOT/f'presentations/{k}.pptx');slides+=len(p.slides)
        if len(p.slides)!=len(DECKS[k]['slides'])+1:errors.append(k+': unexpected presentation count')
        for n,sl in enumerate(p.slides,1):
            if not sl.notes_slide.notes_text_frame.text.strip():errors.append(k+f': empty notes {n}')
            if n<=len(DECKS[k]['slides']):
                for r in DECKS[k]['slides'][n-1]['refs']:
                    if r not in sl.notes_slide.notes_text_frame.text:errors.append(k+f': missing note source {n}')
            for sh in sl.shapes:
                if min(sh.left,sh.top)<0 or sh.left+sh.width>p.slide_width+100 or sh.top+sh.height>p.slide_height+100:errors.append(k+f': bounds {n}')
    for path in publication_files():
        if path.suffix!='.html':continue
        soup=BeautifulSoup(path.read_text(),'html.parser');pageids=[x['id'] for x in soup.select('[id]')]
        if len(pageids)!=len(set(pageids)):errors.append(str(path)+': duplicate HTML ID')
        for a in soup.select('a[href]'):
            u=urlsplit(a['href'])
            if u.scheme or u.netloc:continue
            target=(path.parent/unquote(u.path)).resolve() if u.path else path;links+=1
            if not target.exists():errors.append(str(path.relative_to(ROOT))+': missing '+a['href'])
            elif u.fragment and target.suffix=='.html' and not BeautifulSoup(target.read_text(),'html.parser').find(id=unquote(u.fragment)):errors.append('Missing anchor '+a['href'])
    for folder in ('archive','dist','content','demo-assets'):
        if (ROOT/folder).exists():errors.append('Retained superseded directory '+folder)
    for folder in ('devops','security','finops','architecture','platforms','mechanisms'):
        for p in (ROOT/folder).glob('*'):
            if p.is_file() and (p.suffix!='.html' or len(BeautifulSoup(p.read_text(),'html.parser').get_text().split())>100):errors.append('Competing document edition '+str(p.relative_to(ROOT)))
    if {p.stem for p in (ROOT/'presentations').glob('*.pptx')}!=set(DECKS):errors.append('Unexpected deck catalog')
    x=json.loads((ROOT/'examples/prepared-evidence.json').read_text());f=x['finops']
    assert f['rates_are_not_aws_prices'] and f['same_account_service_and_cost_basis']
    b=c=D('0')
    for row in f['rows']:
        assert D(row['baseline_quantity'])*D(row['unit_rate'])==D(row['baseline_cost'])
        assert D(row['comparison_quantity'])*D(row['unit_rate'])==D(row['comparison_cost'])
        b+=D(row['baseline_cost']);c+=D(row['comparison_cost'])
    assert (b,c,c-b,(c-b)/b*100)==tuple(D(f[a]) for a in ('expected_baseline','expected_comparison','expected_delta','expected_percent_change'))
    assert D(f['prepared_job_log_additional_transfer_units'])*D(f['rows'][2]['unit_rate'])==D('274')
    soup=BeautifulSoup(MD.render(g['finops']['raw']),'html.parser')
    table=next(t for t in soup.select('table') if t.select_one('th') and t.select_one('th').get_text()=='Component')
    for tr,row in zip(table.select('tbody tr')[:3],f['rows']):
        vals=[td.get_text().replace(',','').strip() for td in tr.select('td')]
        assert vals==[row[a] for a in ('component','baseline_quantity','comparison_quantity','unit_rate','baseline_cost','comparison_cost')]
    fixture_results.append({'fixture':'finops','checked':'quantities × rates, component totals, 574 delta, 41% change, transfer-only corroboration and guide-table consistency'})
    dv=x['devops'];assert dv['prior']['database_connections']==dv['prior']['database_limit'];assert dv['current']['database_connections']<dv['current']['database_limit'];assert dv['current']['inventory_p95_ms']>dv['prior']['inventory_p95_ms'];assert dv['current']['database_query_p95_ms']<dv['prior']['database_query_p95_ms']
    fixture_results.append({'fixture':'devops','checked':'contrary current database evidence and elevated upstream latency; no recall or diagnosis execution asserted'})
    se=x['security'];assert se['proposed_revision']==se['deployed_revision'];assert se['test_identity']==se['prepared_revalidation_identity'];assert se['original_revision']!=se['deployed_revision']
    fixture_results.append({'fixture':'security','checked':'prepared revision and identity consistency; no live revalidation asserted'})
    preservation={};excluded_hits=[]
    if (ROOT/'.git').exists():
        changed=git('diff','--name-only','HEAD','--','applications','.github/workflows/validation-console.yml')
        if changed:errors.append('Unrelated application/workflow modified')
        preservation={'checked_against_checkout':git('rev-parse','HEAD'),'application_tree':git('rev-parse','HEAD:applications'),'console_workflow_blob':git('rev-parse','HEAD:.github/workflows/validation-console.yml'),'uncommitted_changes':changed.splitlines()}
        prior='c1c810b48e2349d3fb2479402c767a2834c76b9a'
        old=json.loads(git('show',prior+':content/narratives.json'))['agents']
        terms={v['title'].casefold() for k,v in old.items() if k not in DECKS}
        e=json.loads(git('show',prior+':content/devops-evidence.json'))
        for cap in e.get('capabilities',[]):
            product=cap.get('product','')
            if product and product not in {'AWS DevOps Agent','AWS Security Agent','AWS FinOps Agent'}:terms.add(product.casefold());terms.add(product.removesuffix(' Developer').casefold())
        for p in publication_files():
            rel=p.relative_to(ROOT).as_posix()
            if p.suffix in {'.zip','.bundle','.b64','.gz','.tar'}:errors.append('Stored distribution '+rel)
            text=''
            if p.suffix=='.pptx':
                with ZipFile(p) as z:text=' '.join(z.read(n).decode() for n in z.namelist() if n.endswith(('.xml','.rels')))
            else:
                try:text=p.read_text()
                except UnicodeDecodeError:continue
            if any(t and t in (rel+' '+text).casefold() for t in terms):excluded_hits.append(rel)
        errors.extend('Excluded-scope content '+s for s in excluded_hits)
    if browser:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            br=pw.chromium.launch(headless=True,executable_path=shutil.which('chromium') or None,args=['--no-sandbox'])
            for key in (*DECKS,'index'):
                path=ROOT/('index.html' if key=='index' else f'guides/{key}.html')
                for width in (390,1440):
                    page=br.new_page(viewport={'width':width,'height':1000});js_errors=[];page.on('pageerror',lambda e:js_errors.append(str(e)))
                    page.goto(path.as_uri(),wait_until='load')
                    assert not page.evaluate('document.documentElement.scrollWidth>innerWidth+2'),(key,width,'page overflow')
                    if key!='index':
                        n=page.locator('.report-section').count();page.locator('#search').fill('not-found-qz987');assert page.locator('.report-section:visible').count()==0
                        page.locator('#reset').click();assert page.locator('.report-section:visible').count()==n
                        if key!='architecture':page.locator('#kind').select_option('mechanism');assert page.locator('.report-section:visible').count()>0
                        page.locator('#reset').click();sid=next(iter(g[key]['sections']));page.locator('#search').fill('not-found-qz987');page.evaluate('(id)=>location.hash=id',sid);page.wait_for_timeout(100);assert page.locator('#'+sid).is_visible()
                        page.locator('#present').click();assert page.locator('.slide:visible').count()==1
                        page.locator('#next').click();assert page.locator('#slide-select').input_value()=='1'
                        page.locator('#notes-toggle').check();assert page.locator('.notes:visible').count()==1
                        page.locator('#notes-toggle').uncheck()
                        for i in range(len(DECKS[key]['slides'])):
                            page.locator('#slide-select').select_option(str(i));assert page.locator('.slide:visible').count()==1
                            assert not page.evaluate('document.documentElement.scrollWidth>innerWidth+2'),(key,width,i,'slide overflow')
                        page.locator('#next').focus();page.keyboard.press('Home');assert page.locator('#slide-select').input_value()=='0'
                        page.locator('#timer-toggle').click();page.wait_for_timeout(1150);page.locator('#timer-toggle').click();assert page.locator('#timer').inner_text()!='0:00'
                        page.locator('#timer-reset').click();page.wait_for_timeout(300);assert page.locator('#timer').inner_text()=='0:00'
                        page.emulate_media(media='print');assert page.locator('.slide:visible').count()==len(DECKS[key]['slides']);page.emulate_media(media='screen');assert page.locator('.slide:visible').count()==1
                        page.keyboard.press('Escape');assert page.locator('.shell').is_visible()
                    errors.extend(js_errors);browser_results.append({'page':key,'width':width,'javascript_errors':js_errors,'horizontal_overflow':False,'interaction_checks':key!='index'});page.close()
            br.close()
    result={'revision':REV,'canonical_guides':len(g),'canonical_sections':sectioncount,'unique_official_urls':len(source_urls),'local_links_checked':links,'timed_slides':sum(len(d['slides']) for d in DECKS.values()),'pptx_slides_including_appendices':slides,'programme_minutes':sum(s['seconds'] for d in DECKS.values() for s in d['slides'])//60,'exercise_minutes':sum(s['seconds'] for d in DECKS.values() for s in d['slides'] if s['kind']=='demo')//60,'synthetic_fixture_checks':fixture_results,'browser':browser_results,'preservation':preservation,'excluded_scope_hits':excluded_hits,'errors':errors,'boundary':'Checks cover document structure, local links, source identifiers, synthetic arithmetic/record consistency, browser behavior, slide bounds and notes. They are not external source-content verification, successful AWS configuration, causal inference testing or measured operational outcomes.'}
    write('qa/checks.json',json.dumps(result,indent=2)+'\n')
    if errors:raise SystemExit(json.dumps(result,indent=2))
    write('VALIDATION.md','# Validation and evidence boundary\n\nRevision: '+REV+'.\n\nThe completed automated checks are recorded in [qa/checks.json](qa/checks.json). They cover '+str(sectioncount)+' globally identified canonical sections, '+str(links)+' local HTML links, '+str(slides)+' PowerPoint slides, note references and the explicit synthetic fixtures. Browser checks run at 390 and 1440 pixels when `--browser` is used; the JSON records whether they actually ran.\n\nThe source review is editorial work, separate from the automated checks. Each guide identifies its reviewed official sections, unresolved or conflicting source versions, and detailed versus linked/deferred coverage. A passing reference-ID check does not prove a source supports every claim. See [content review](qa/content-review.md).\n\nThe new presentation programme is 124 minutes including 28 minutes of prepared-result exercises. The evidence appendix in each deck is untimed. These are walkthrough budgets, not live service completion guarantees.\n\nPresentation layout was rendered locally from the curated slide source and representative slides inspected; full canonical reference text was subsequently embedded into notes by the publication build. Automated bounds checks are not pixel-level assurance of every possible office renderer.\n\nNo live AWS workflow, credential provider, security test, memory-recall experiment or external destination publication was executed. No improvement, security coverage or saving is measured. The readiness-console application and its workflow are outside the documentation build and are checked for preservation.\n\nLegacy HTML paths contain only navigation stubs to the canonical guides. They are not retained document editions. Git history holds prior content; no manual archives or ZIP distributions are generated.\n')
    files={p.relative_to(ROOT).as_posix():{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in publication_files()}
    write('MANIFEST.json',json.dumps({'revision':REV,'scope':'Documentation files only; application has its own publication records','files':files,'excluded':['MANIFEST.json itself','applications/','.git/']},indent=2)+'\n')
    print(json.dumps(result,indent=2))
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--browser',action='store_true');a=p.parse_args();check(a.browser)
