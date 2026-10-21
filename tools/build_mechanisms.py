#!/usr/bin/env python3
"""Build source-linked mechanism companions without changing core presentation timing."""
from __future__ import annotations
from pathlib import Path
from html import escape as H
import json
from datetime import datetime

CSS=r'''
:root{--ink:#182d43;--muted:#52657a;--line:#d9e3ec;--paper:#fff;--bg:#f3f6f9;--accent:#176e83;--navy:#10253b}*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:84px}body{overflow-wrap:anywhere;margin:0;background:var(--bg);color:var(--ink);font:16px/1.65 system-ui,-apple-system,Segoe UI,sans-serif}a{color:var(--accent);text-underline-offset:3px}a:hover{text-decoration-thickness:2px}button,input,select{font:inherit}button,.button{display:inline-block;border:1px solid var(--line);border-radius:7px;padding:9px 14px;background:var(--paper);color:var(--ink);cursor:pointer;text-decoration:none}button:hover,.button:hover{border-color:var(--accent)}button:focus-visible,input:focus-visible,select:focus-visible,a:focus-visible{outline:3px solid #3182ce;outline-offset:3px}h1,h2,h3,h4{line-height:1.2;letter-spacing:-.02em}h1{font-size:clamp(34px,4vw,54px);max-width:980px;margin:16px 0}h2{font-size:30px;margin:12px 0 20px}h3{font-size:22px}h4{font-size:17px;margin:16px 0 9px}p{margin:10px 0 16px}.wrap{max-width:1190px;margin:auto;padding:0 30px}.hero{background:var(--navy);color:white;padding:40px 0 44px}.hero a{color:#d3e5f6}.eyebrow{font-size:12px;letter-spacing:.13em;text-transform:uppercase;font-weight:750}.lead{font-size:20px;max-width:950px;color:#d3deea}.meta,.actions{display:flex;gap:10px;flex-wrap:wrap;margin-top:20px}.badge{display:inline-block;padding:4px 10px;font-size:12px;border:1px solid var(--line);border-radius:4px;background:#f4f7fb;color:var(--ink);font-weight:650}.hero .badge{color:#dfe9f3;border-color:#42566a;background:#1a334b}.hero .button,.hero button{color:var(--ink)}.nav{position:sticky;top:0;background:#fffefc;border-bottom:1px solid var(--line);z-index:9}.nav .wrap{display:flex;gap:22px;overflow:auto;padding-top:12px;padding-bottom:12px;white-space:nowrap}.nav a{font-size:14px;color:var(--ink)}section{margin:40px 0 60px}.panel,.record,.stage,.source{background:var(--paper);border:1px solid var(--line);border-radius:10px;padding:24px}.grid{display:grid;grid-template-columns:1fr 1fr;gap:20px}.notice{padding:18px 22px;border-left:4px solid var(--accent);background:#eaf0f6;margin:20px 0}.small{font-size:13px;color:var(--muted)}.banner{padding:15px 20px;background:#fff4db;border:1px solid #e4cca0;border-radius:8px;margin:20px 0}.flow{padding:24px;border-radius:10px;background:var(--navy);color:white;font-size:21px;line-height:1.65}.toolbar{display:flex;align-items:end;flex-wrap:wrap;gap:14px;padding:20px;background:white;border:1px solid var(--line);border-radius:10px;margin:20px 0}.toolbar label{display:grid;gap:5px;font-size:13px;font-weight:650}.toolbar input,.toolbar select{border:1px solid #a9bacb;border-radius:5px;min-height:43px;padding:8px 12px;background:#fff;color:var(--ink)}.toolbar input{width:310px}.count{font-variant-numeric:tabular-nums;margin-left:auto}.record{padding:0;margin:18px 0;overflow:hidden}.record summary{padding:23px 25px;cursor:pointer;list-style:none;display:flex;justify-content:space-between;gap:18px;align-items:start}.record summary::-webkit-details-marker{display:none}.record summary:after{content:'+';font-size:24px;line-height:1}.record[open] summary:after{content:'−'}.record h3{margin:5px 0 12px}.record .body{border-top:1px solid var(--line);padding:0 25px 24px}.record .function{font-size:17px;max-width:1000px}.label{font-size:11px;letter-spacing:.08em;font-weight:750;text-transform:uppercase;color:var(--muted);margin-top:22px}.body ol,.body ul{padding-left:22px}.body li{margin:8px 0}.body .split{display:grid;grid-template-columns:1fr 1fr;gap:26px}.body .box{background:var(--bg);border-radius:7px;padding:16px 20px}.body .limit{border-left:3px solid #9c6817;padding-left:17px}.refs{font-size:13px;margin:16px 0}.refs a{margin-right:15px;display:inline-block}.stepnav{display:flex;flex-wrap:wrap;gap:9px;margin:20px 0}.stepnav button[aria-pressed=true]{color:white;background:var(--accent)}.stage{margin:16px 0;border-left:4px solid var(--accent)}.stage h3{margin:6px 0 14px}.stage strong{font-size:14px}.checklist label{display:flex;align-items:start;gap:12px;margin:15px 0}.checklist input{margin-top:7px;width:18px;height:18px;flex-shrink:0}.source{padding:18px 22px;margin:12px 0;overflow-wrap:anywhere}.source h3{font-size:18px;margin:0 0 8px}.tablewrap{overflow:auto}table{border-collapse:collapse;width:100%;background:white;font-size:14px}th,td{padding:15px;border:1px solid var(--line);text-align:left;vertical-align:top}th{background:var(--navy);color:#fff}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:14px/1.6 ui-monospace,monospace;background:#eef3f8;padding:18px;border-radius:7px}.footer{border-top:1px solid var(--line);padding:30px 0;color:var(--muted);font-size:13px}.slidepanel{background:#fff;border:1px solid var(--line);border-radius:12px;min-height:540px;padding:35px}.slidepanel h3{font-size:30px;margin:10px 0 25px}.slidepanel .grid .panel{background:#f4f7fa}.slidepanel h4{color:var(--accent);font-size:14px;text-transform:uppercase;letter-spacing:.08em}.slidepanel .content{font-size:20px}.slidecontrols{display:flex;gap:10px;flex-wrap:wrap;align-items:center;margin:15px 0}.slidecount{margin-left:auto}.fullscreen{position:fixed;inset:0;z-index:20;overflow:auto;background:var(--bg);padding:22px 4vw;margin:0}.fullscreen .slidepanel{max-width:1440px;margin:auto}.fullscreen .slidecontrols{max-width:1440px;margin:15px auto}.fullscreen h2{display:none}.hide{display:none!important}[hidden]{display:none!important}.skip{position:absolute;left:-999px}.skip:focus{left:10px;top:10px;background:white;padding:10px;z-index:99}.toast{position:fixed;bottom:20px;right:20px;background:var(--navy);color:white;padding:12px 18px;border-radius:6px;z-index:30}
@media(max-width:760px){.wrap{padding:0 18px}.grid,.body .split{grid-template-columns:1fr}.hero{padding:28px 0}.lead{font-size:18px}section{margin:30px 0 44px}.toolbar input{width:100%}.toolbar label{width:100%}.count{margin-left:0}.record summary{padding:20px}.record .body{padding:0 20px 20px}.slidepanel{padding:22px;min-height:auto}.slidepanel h3{font-size:25px}.slidepanel .content{font-size:17px}h2{font-size:26px}.flow{font-size:18px}}
@media print{body{background:white;color:#111;font-size:11pt}.hero{background:white!important;color:#111;padding:0}.hero a,.lead{color:#333}.nav,.toolbar,.actions,.stepnav,.checklist button,.slidecontrols,.toast,#reference{display:none!important}.wrap{max-width:none;padding:0}section{margin:20px 0}.record,.stage,.panel{break-inside:avoid}.record[hidden],.stage[hidden]{display:block!important}.grid{display:block}.footer{padding:12px 0}.source{padding:9px}.hero .badge{color:#111;border-color:#bbb}.flow{color:#111;background:#eee}}
'''
JS=r'''
(()=>{const q=s=>document.querySelector(s), all=s=>[...document.querySelectorAll(s)];
const inventory=JSON.parse(q('#mechanism-data').textContent), records=all('.record');
function filter(){const text=(q('#search')?.value||'').toLowerCase(), family=q('#family')?.value||'';let n=0;records.forEach(r=>{r.hidden=!(r.textContent.toLowerCase().includes(text)&&(!family||r.dataset.family===family));if(!r.hidden)n++});if(q('#count'))q('#count').textContent=`${n} / ${records.length} mechanisms`;}
q('#search')?.addEventListener('input',filter);q('#family')?.addEventListener('change',filter);q('#reset')?.addEventListener('click',()=>{q('#search').value='';q('#family').value='';filter()});
q('#expand')?.addEventListener('click',()=>records.filter(r=>!r.hidden).forEach(r=>r.open=true));q('#collapse')?.addEventListener('click',()=>records.forEach(r=>r.open=false));
function save(name,data){const url=URL.createObjectURL(new Blob([JSON.stringify(data,null,2)],{type:'application/json'}));const a=document.createElement('a');a.href=url;a.download=name;a.click();setTimeout(()=>URL.revokeObjectURL(url),1000)}
q('#export')?.addEventListener('click',()=>save('mechanism-reference.json',inventory));
all('[data-step]').forEach(b=>b.addEventListener('click',()=>{all('[data-step]').forEach(x=>x.setAttribute('aria-pressed',String(x===b)));all('.stage').forEach(x=>x.hidden=b.dataset.step!=='all'&&x.dataset.stage!==b.dataset.step)}));
q('#save-checks')?.addEventListener('click',()=>save('mechanism-review-worksheet.json',{type:'Manual proposed checks; not test results or authorization',created:new Date().toISOString(),agent:document.body.dataset.agent,checks:all('.checklist input').map(x=>({text:x.dataset.text,checked:x.checked}))}));
q('#clear-checks')?.addEventListener('click',()=>all('.checklist input').forEach(x=>x.checked=false));q('#print')?.addEventListener('click',()=>window.print());
let saved;window.addEventListener('beforeprint',()=>{saved=records.map(r=>r.open);records.forEach(r=>r.open=true)});window.addEventListener('afterprint',()=>{records.forEach((r,i)=>r.open=saved?.[i]||false)});
const slides=all('.reference-slide');let current=0;function show(n){if(!slides.length)return;current=Math.max(0,Math.min(n,slides.length-1));slides.forEach((s,i)=>s.hidden=i!==current);q('#slide-count').textContent=`Reference ${current+1} / ${slides.length} · untimed`;q('#prev').disabled=current===0;q('#next').disabled=current===slides.length-1}
q('#prev')?.addEventListener('click',()=>show(current-1));q('#next')?.addEventListener('click',()=>show(current+1));q('#focus')?.addEventListener('click',()=>{q('#reference').classList.toggle('fullscreen');q('#focus').textContent=q('#reference').classList.contains('fullscreen')?'Exit presentation view':'Presentation view'});
document.addEventListener('keydown',e=>{if(!q('#reference')?.classList.contains('fullscreen')||/INPUT|SELECT|TEXTAREA/.test(e.target.tagName))return;if(e.key==='ArrowRight'){show(current+1);e.preventDefault()}if(e.key==='ArrowLeft'){show(current-1);e.preventDefault()}if(e.key==='Escape'){q('#reference').classList.remove('fullscreen');q('#focus').textContent='Presentation view'}});show(0);filter();
})();
'''

def p(s):return f'<p>{H(str(s))}</p>'
def ul(items,ordered=False):
    tag='ol' if ordered else 'ul';return '<'+tag+'>'+''.join('<li>'+H(str(s))+'</li>' for s in items)+'</'+tag+'>'
def refs(ids,sources,prefix='src-'):
    return '<div class="refs">Official guidance: '+''.join(f'<a href="#{prefix}{H(k)}">[{H(k)}] {H(sources[k]["title"])}</a>' for k in ids)+'</div>'
def stamp(s):return f'{s//60:02}:{s%60:02}'

def record_html(r,sources,prefix=''):
    return f'''<details class="record" id="{prefix}{r['id']}" data-family="{H(r['family'])}"><summary><div><span class="eyebrow">{r['id']} / {H(r['family'])}</span><h3>{H(r['title'])}</h3><span class="badge">{H(r['surface'])}</span> <span class="badge">{H(r['activation'])}</span></div></summary><div class="body"><div class="label">01 / Identity and availability</div>{p(r['availability'])}<div class="label">02 / Documented mechanism</div><div class="function">{p(r['mechanism'])}</div><div class="split"><div><div class="label">03 / Enable and configure — official procedure</div><p><strong>Prerequisites:</strong> {H(r['prerequisites'])}</p>{ul(r['steps'],True)}<div class="label">Scope and persistence</div>{p(r['scope'])}</div><div><div class="label">04 / Independently verify — proposed team exercise</div>{ul(r['checks'])}<div class="label">05 / Record to retain — proposed</div>{p(r['artifact'])}<div class="label">06 / Stop, correct or remove</div>{p(r['control'])}</div></div><div class="limit"><h4>Limitations and effects</h4>{p(r['limits'])}</div><div class="box"><h4>Illustrative task — not executed</h4>{p(r['example'])}<p><strong>Fallback:</strong> {H(r['fallback'])}</p></div><div class="label">07 / Evidence reviewed 21 October 2026</div>{refs(r['refs'],sources,prefix+'src-')}</div></details>'''

def source_html(ids,sources,prefix=''):
    out='<section id="'+prefix+'sources"><div class="eyebrow">Official documentation</div><h2>Evidence register</h2><p class="small">Links support the documented mechanisms. Independent checks, retained records and example prompts are proposed engineering practices. Page content can change after the review date.</p>'
    for k in ids:
        s=sources[k];out+=f'<article class="source" id="{prefix}src-{k}"><h3>[{k}] {H(s["title"])}</h3><a href="{H(s["url"])}" target="_blank" rel="noopener noreferrer">{H(s["url"])}</a><p class="small">Reviewed {s["reviewed"]}. {H(s["note"])}</p></article>'
    return out+'</section>'

def usecase_html(a,sources):
    u=a['use_case'];out=f'<section id="usecase"><div class="eyebrow">Proposed use-case exercise</div><h2>{H(u["title"])}</h2><p>Optional {u["minutes"]}-minute prepared-result walkthrough. This is additional to the existing core session; agent processing is not guaranteed to complete within the allocation.</p><div class="notice"><strong>Preflight</strong>{p(u["preflight"])}</div><div class="stepnav"><button data-step="all" aria-pressed="true">All steps</button>'
    for i in range(6):out+=f'<button data-step="{i}" aria-pressed="false">{i+1}</button>'
    out+='</div>';c=0
    for i,(title,body,role,gate,rid,seconds) in enumerate(u['steps']):
        out+=f'<article class="stage" data-stage="{i}"><div class="eyebrow">Step {i+1} / {stamp(c)}–{stamp(c+seconds)} / {H(role)}</div><h3>{H(title)}</h3>{p(body)}<p><strong>Acceptance gate:</strong> {H(gate)}</p><a href="#{rid}">Mechanism {rid}</a></article>';c+=seconds
    return out+f'<div class="grid"><div class="panel"><h3>Exit criterion</h3>{p(u["exit"])}</div><div class="panel"><h3>Prepared-result fallback</h3>{p(u["fallback"])}</div></div>{refs(u["refs"],sources)}</section>'

def reference_html(a,sources):
    out='<section id="reference"><h2>Mechanism reference presentation</h2><p class="small">Untimed discussion material. Full setup instructions and evidence remain in each record and in the PowerPoint notes.</p><div class="slidecontrols"><button id="prev">Previous</button><button id="next">Next</button><button id="focus">Presentation view</button><span id="slide-count" class="slidecount" role="status"></span></div>'
    for r in a['records']:
        out+=f'<article class="reference-slide slidepanel" hidden><div class="eyebrow">{r["id"]} / {H(r["family"])} / {H(a["title"])}</div><h3>{H(r["title"])}</h3><div class="grid"><div class="panel"><h4>Documented operation</h4><p class="content">{H(r["mechanism"])}</p><h4>Activate</h4>{p(r["slide"][1])}</div><div class="panel"><h4>Independently verify — proposed</h4><p class="content">{H(r["checks"][-1])}</p><h4>Boundary</h4>{p(r["limits"])}</div></div><p class="small">{H(r["surface"])} · {H(r["availability"])}</p><details><summary>Setup and speaker detail</summary>{ul(r["steps"],True)}{p(r["control"])}{refs(r["refs"],sources)}</details></article>'
    return out+'</section>'

def shell(title,lead,body,key='overview',accent='#176e83',data=None,subtitle='Mechanisms and enablement'):
    payload=json.dumps(data or {},ensure_ascii=False).replace('<','\\u003c')
    return f'''<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="description" content="Source-linked mechanism, setup and independent verification reference for {H(title)}"><title>{H(title)} | Mechanisms and enablement</title><style>{CSS}</style></head><body data-agent="{key}" style="--accent:{accent}"><a class="skip" href="#overview">Skip to content</a><header class="hero"><div class="wrap"><a href="../index.html">← Report collection</a><div class="eyebrow" style="margin-top:25px">{H(subtitle)} / Engineering reference</div><h1>{H(title)}</h1><p class="lead">{H(lead)}</p><div class="meta"><span class="badge">Reviewed 21 October 2026</span><span class="badge">Documented behavior ≠ measured outcome</span><span class="badge">No live AWS execution</span></div><div class="actions"><a class="button" href="index.html">Cross-platform index</a><a class="button" href="{key}.md" download>Markdown reference</a><a class="button" href="../presentations/{key}-mechanisms-reference.pptx" download>PowerPoint reference</a><button id="print">Print guide</button></div></div></header><nav class="nav" aria-label="Guide sections"><div class="wrap"><a href="#overview">Mechanism model</a><a href="#inventory">Inventory</a><a href="#usecase">Use case</a><a href="#reference">Reference slides</a><a href="#review">Review worksheet</a><a href="#sources">Official sources</a></div></nav><main class="wrap">{body}</main><footer class="footer"><div class="wrap">Independent documentation synthesis. Not an AWS publication, a deployed orchestrator, an approval or a certification. <a href="../EDITORIAL_STANDARD.md">Editorial standard</a> · <a href="VALIDATION.md">Mechanism validation boundary</a></div></footer><script type="application/json" id="mechanism-data">{payload}</script><script>{JS}</script></body></html>'''

def checklist():
    items=['Product, surface, version and availability checked.','Data and target authorization confirmed; secrets excluded from retained context.','Official activation steps followed and resulting setting recorded.','Independent evidence collected in a new session or job where appropriate.','Stop, correction and residual-data behavior checked.','Owner accepted the result or recorded the unresolved limitation.']
    return '<section id="review"><div class="eyebrow">Proposed manual worksheet</div><h2>Record what was actually checked</h2><div class="panel checklist"><p>These selections are browser-local form state. Exported entries are self-reported, not automated test evidence or authorization.</p>'+''.join(f'<label><input type="checkbox" data-text="{H(x)}"><span>{H(x)}</span></label>' for x in items)+'<div class="actions"><button id="save-checks">Export worksheet JSON</button><button id="clear-checks">Clear</button></div></div></section>'

def guide_html(key,a,sources):
    rs=a['records'];ids=list(dict.fromkeys(a['learning_refs']+[k for r in rs for k in r['refs']]+a['use_case']['refs']))
    body=f'<section id="overview"><div class="eyebrow">Context, configuration and control</div><h2>What the platform retains and how it is used</h2>{p(a["opening"])}<div class="flow">{H(a["learning"])}</div>{refs(a["learning_refs"],sources)}<div class="banner"><strong>Availability and scope</strong>{p(a["status"])}</div><p class="small">This guide extends the <a href="../{a["report"]}">capability report</a>. It does not replace its historical source records or change its timed presentation.</p></section>'
    families=sorted(set(r['family'] for r in rs))
    body+=f'<section id="inventory"><div class="eyebrow">{len(rs)} mechanism records</div><h2>Documented operation → setup → independent verification</h2><div class="toolbar"><label>Search records<input id="search" type="search" placeholder="Memory, permissions, setup…"></label><label>Mechanism family<select id="family"><option value="">All families</option>'+''.join(f'<option>{H(f)}</option>' for f in families)+'</select></label><button id="reset">Reset</button><span id="count" class="small count" role="status"></span><button id="expand">Expand visible</button><button id="collapse">Collapse all</button><button id="export">Export inventory JSON</button></div>'+''.join(record_html(r,sources) for r in rs)+'</section>'
    body+=usecase_html(a,sources)+reference_html(a,sources)+checklist()+source_html(ids,sources)
    return shell(a['title'],a['opening'],body,key,a['accent'],{'agent':key,'records':rs,'sources':{k:sources[k] for k in ids}})

def guide_md(key,a,sources):
    lines=[f'# {a["title"]} — mechanisms and enablement','Reviewed 21 October 2026. Independent checks and examples are proposed; no live workflow was executed.','## Mechanism model',a['opening'],a['learning'],a['status']]
    for r in a['records']:
        lines += [f'## {r["id"]} — {r["title"]}',f'**Surface:** {r["surface"]}\n\n**Availability:** {r["availability"]}\n\n**Activation:** {r["activation"]}','### Documented mechanism',r['mechanism'],'### Enable and configure',r['prerequisites']]
        lines += [f'{i+1}. {x}' for i,x in enumerate(r['steps'])]
        lines += ['### Scope and persistence',r['scope'],'### Independent checks — proposed']+[f'- {x}' for x in r['checks']]
        lines += ['### Record to retain — proposed',r['artifact'],'### Stop, correct or remove',r['control'],'### Limits and fallback',r['limits'],r['fallback'],'### Illustrative task — not executed',r['example'],'### Official guidance']
        lines += [f'[{k}: {sources[k]["title"]}]({sources[k]["url"]})' for k in r['refs']]
    u=a['use_case'];lines += ['## Proposed use-case exercise',u['title'],f'Optional {u["minutes"]}-minute walkthrough; additional to the core programme. Not an agent-runtime guarantee.',u['preflight']];c=0
    for i,(title,body,role,gate,rid,sec) in enumerate(u['steps']):
        lines += [f'### {i+1}. {title} · {stamp(c)}–{stamp(c+sec)}',f'**Role:** {role} · **Mechanism:** {rid}',body,f'**Gate:** {gate}'];c+=sec
    lines += ['### Exit',u['exit'],'### Fallback',u['fallback']]
    return '\n\n'.join(lines)+'\n'

def overview_html(data):
    sources=data['sources'];body='<section id="overview"><div class="eyebrow">Cross-platform comparison</div><h2>“Learning” must name a mechanism</h2><p>Retained knowledge, retrieved documents, maintained instructions, conversation summarization and model training are different operations. This collection documents controls visible in official guidance; it does not infer undisclosed internals.</p><div class="notice"><strong>How to read this extension</strong><p>The existing five presentations remain 160 minutes including 48 minutes of demonstrations. New mechanism slides are untimed reference material. Optional exercises add 44 minutes when delivered separately; they are not silently included in the original programme.</p></div></section><section id="inventory"><h2>Choose a platform-specific guide</h2><div class="grid">'
    for k,a in data['agents'].items():
        body+=f'<article class="panel"><div class="eyebrow">{len(a["records"])} mechanisms / {a["use_case"]["minutes"]}-minute optional exercise</div><h3>{H(a["title"])}</h3>{p(a["learning"])}<a class="button" href="{k}.html">Open mechanism guide</a></article>'
    body+='</div><h3>Mechanism and activation comparison</h3><div class="tablewrap"><table><thead><tr>'+''.join('<th>'+H(x)+'</th>' for x in ['Platform','Retained or supplied context','Activation','Configuration surface','Independent evidence','Limit'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+H(x)+'</td>' for x in row)+'</tr>' for row in data['comparison'])+'</tbody></table></div></section><section id="usecase"><h2>A common review sequence, not a shared runtime</h2><div class="flow">Identify the mechanism → authorize the scope → enable deliberately → observe a new session/job → challenge the result → correct or stop</div><p>Apply this sequence through each product’s own controls. A handoff transfers evidence and scope, not credentials, remembered state or approval authority.</p></section><section id="reference"><h2>Cross-platform operating controls</h2><p class="small">The following six controls are proposed team practices, supported by product-specific evidence.</p>'
    for i,(title,text,ids) in enumerate(data['controls']):body+=f'<article class="stage"><div class="eyebrow">Control {i+1}</div><h3>{H(title)}</h3>{p(text)}{refs(ids,sources)}</article>'
    body+='</section>'+checklist()+'<section><h2>Evidence qualifications and changes</h2>'+ul(data['evidence_notes'])+'</section>'+source_html(list(sources),sources)
    return shell('Platform mechanisms and enablement','A source-linked reference for retained context, instructions, triggers, tools, permissions and independent verification.',body,'overview',data=data)

# PowerPoint: editable shapes; full procedures and source URLs in speaker notes.
def pptx_refs(key,a,sources,out):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    navy='10253B';ink='182D43';muted='52657A';accent=a.get('accent','#176e83').lstrip('#');white='FFFFFF';soft='F1F5F9'
    def shape(s,x,y,w,h,color):
        z=s.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h));z.fill.solid();z.fill.fore_color.rgb=RGBColor.from_string(color);z.line.fill.background();return z
    def text(s,txt,x,y,w,h,sz=19,color=ink,bold=False):
        t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=t.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
        for i,line in enumerate(str(txt).split('\n')):
            p0=tf.paragraphs[0] if i==0 else tf.add_paragraph();p0.text=line;p0.font.name='Aptos';p0.font.size=Pt(sz);p0.font.bold=bold;p0.font.color.rgb=RGBColor.from_string(color);p0.space_after=Pt(6)
        return t
    def base(title,kicker,notes):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(white)
        shape(s,0,0,13.333,.14,accent);text(s,kicker.upper(),.55,.38,12.2,.28,11,accent,True)
        size=29 if len(title)<72 else 26;text(s,title,.55,.85,12.2,.95,size,ink,True)
        shape(s,.55,7.08,12.2,.012,'D4DFE9');text(s,a['title']+' | Mechanism reference | 21 Oct 2026',.55,7.18,10.5,.2,9,muted);text(s,f'{len(prs.slides):02} / UNTImED'.upper(),11.1,7.18,1.7,.2,9,muted)
        s.notes_slide.notes_text_frame.text=notes
        return s
    def box(s,label,body,x,y,w,h,sz=18):
        shape(s,x,y,w,h,soft);text(s,label.upper(),x+.2,y+.15,w-.4,.25,11,accent,True);text(s,body,x+.2,y+.55,w-.4,h-.67,sz)
    s=base(a['title']+' — mechanisms and enablement','Reference companion',a['opening']+'\n\n'+a['learning']+'\n\n'+data_boundary())
    box(s,'How retained context is used',a['learning'],.55,2.05,7.45,3.05,23);box(s,'Reading contract','Documented operation\nOfficial setup\nProposed independent checks\nStop and residual-data limits',8.25,2.05,4.53,3.05,21)
    text(s,'Untimed reference slides. The existing core programme and demonstration allocations remain unchanged.',.65,5.5,11.9,.8,22,muted)
    transcripts=['# '+a['title']+' — mechanism reference speaker notes','All reference slides are untimed. Optional exercises are additional to the core programme.']
    for r in a['records']:
        notes=guide_md(key,{**a,'records':[r]},sources).split('## Proposed use-case exercise')[0]
        s=base(r['title'],r['id']+' / '+r['family'],notes)
        box(s,'Proposed team control' if key=='overview' else 'Documented operation',r.get('slide_mechanism',r['mechanism']),.55,1.98,6,2.11,18)
        steptext='\n'.join(f'{i+1}. {x}' for i,x in enumerate(r.get('slide_setup',r['steps'])))
        box(s,'Apply through product controls' if key=='overview' else 'Enable / configure',steptext,.55,4.25,6,2.5,20)
        box(s,'Independent check — proposed',r['checks'][-1],6.78,1.98,6,2.11,19)
        box(s,'Control and limitation',r.get('slide_boundary',r['limits']),6.78,4.25,6,2.5,20)
        # Source IDs at bottom; complete URLs are in notes and companion guide.
        text(s,'Official sources: '+', '.join(r['refs'])+' · '+r['surface'],.65,6.82,12,.2,9,muted)
        transcripts += [f'## {r["id"]} — {r["title"]}',r['mechanism'],'**Setup:** '+ ' '.join(r['steps']),'**Proposed verification:** '+ ' '.join(r['checks']),'**Boundary:** '+r['limits']]
    u=a['use_case'];c=0
    for part in range(2):
        subset=u['steps'][part*3:(part+1)*3];notes='PROPOSED EXERCISE. '+u['preflight']+'\nExit: '+u['exit']+'\nFallback: '+u['fallback']
        s=base(u['title'],f'Optional {u["minutes"]}-minute exercise / part {part+1} of 2',notes)
        for j,(title,body,role,gate,rid,sec) in enumerate(subset):
            y=2.0+j*1.53;shape(s,.55,y,12.23,1.35,soft);text(s,f'{stamp(c)}–{stamp(c+sec)}',.76,y+.18,1.45,.33,16,accent,True);text(s,title,2.33,y+.12,10.15,.33,20,ink,True);text(s,body+' Gate: '+gate,2.33,y+.55,10.05,.68,16);c+=sec
    sourceids=list(dict.fromkeys([x for r in a['records'] for x in r['refs']]))
    s=base('Evidence and delivery boundary','Official documentation', '\n'.join(sources[k]['title']+'\n'+sources[k]['url'] for k in sourceids)+'\n\n'+data_boundary())
    text(s,'Official source register',.65,2,12,.5,24,ink,True)
    for j,k in enumerate(sourceids):
        x=.65+(j%2)*6.2;y=2.8+(j//2)*.37
        if y<6.35:
            t=text(s,'['+k+'] '+sources[k]['title'],x,y,5.9,.31,13,accent);t.text_frame.paragraphs[0].runs[0].hyperlink.address=sources[k]['url']
    text(s,'Reviewed 21 October 2026. Proposed exercises were not executed. Consult the full HTML guide for qualifications and source changes.',.65,6.45,12,.48,15,muted)
    if key!='overview':
        transcripts += ['## Optional use-case exercise',guide_md(key,{**a,'records':[]},sources).split('## Proposed use-case exercise')[1]]
    out.parent.mkdir(exist_ok=True);prs.save(out)
    return len(prs.slides),'\n\n'.join(transcripts)+'\n'

def data_boundary():return 'Documented configuration, not a claim about undisclosed internals. No live AWS test or measured improvement. Retained context does not establish customer-controlled model training.'

def overview_ppt(data,out):
    # Reuse the reference layout for six proposed cross-platform controls.
    a={'title':'Common mechanism controls','accent':'#176e83','opening':data_boundary(),'learning':'Identify the scope → select the context → authorize the effect → verify the result → correct or stop. No shared agent runtime or memory store is asserted.','learning_refs':[], 'status':'Proposed architecture practices, not native portfolio functionality.','records':[],'use_case':data['agents']['devops']['use_case']}
    for i,(title,body,ids) in enumerate(data['controls']):
        a['records'].append({'id':f'X-M{i+1:02}','title':title,'family':'Proposed control','surface':'Across the four products; implementation differs','availability':'Proposed team practice','activation':'Apply through each product’s own controls','mechanism':body,'prerequisites':'Use the product-specific setup record and its current official guidance.','steps':['Select the relevant product mechanism.','Record its scope, owner and activation setting.','Run the proposed independent check using approved test inputs.'],'scope':'No shared permission or memory service is asserted.','checks':['Compare the configured boundary with independent evidence from the product or target.'],'artifact':'Product identifier, setting, test evidence and reviewer decision.','control':'Use the product-specific stop or correction control.','limits':'A common review process does not create a common runtime, credential set or approval model.','refs':ids,'fallback':'Use the existing team process when the documented mechanism is unavailable.','example':'Proposed cross-platform review, not a native AWS orchestration configuration.','slide':['Product-specific mechanism','Apply through product controls','Retain independent evidence','No shared runtime implied']})
    # Only the title, six controls and evidence appendix; remove the reused exercise pages.
    n,trans=pptx_refs('overview',a,data['sources'],out)
    from pptx import Presentation
    prs=Presentation(out)
    for idx in [len(prs.slides)-2,len(prs.slides)-3]:
        rid=prs.slides._sldIdLst[idx].rId;prs.part.drop_rel(rid);prs.slides._sldIdLst.remove(prs.slides._sldIdLst[idx])
    
    for i,sl in enumerate(prs.slides,1):
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text.endswith(' / UNTIMED'):
                sh.text_frame.paragraphs[0].runs[0].text=f'{i:02} / UNTIMED'
    prs.save(out);return len(prs.slides),trans

def build_mechanisms(root:Path):
    data=json.loads((root/'content/mechanisms.json').read_text());sources=data['sources'];output=root/'mechanisms';output.mkdir(exist_ok=True)
    def put(rel,txt):p0=root/rel;p0.parent.mkdir(parents=True,exist_ok=True);p0.write_text(txt,encoding='utf-8')
    counts={}
    for key,a in data['agents'].items():
        put(f'mechanisms/{key}.html',guide_html(key,a,sources));put(f'mechanisms/{key}.md',guide_md(key,a,sources))
        n,trans=pptx_refs(key,a,sources,root/f'presentations/{key}-mechanisms-reference.pptx');counts[key]=n;put(f'mechanisms/{key}-presentation-notes.md',trans)
        # Main report retains all existing content; the mechanism section precedes its evidence register.
        path=root/a['report'];text0=path.read_text();prefix='mechanism-'
        inline=f'<section class="section" id="mechanisms"><div class="kicker">Mechanisms / setup and verification extension</div><h2>How this platform uses context and executes work</h2>{p(a["opening"])}<div class="notice">{H(a["learning"])}</div><p><a class="button" href="../mechanisms/{key}.html">Detailed setup guide and optional exercise</a> <a class="button" href="../presentations/{key}-mechanisms-reference.pptx">Untimed reference PowerPoint</a></p><p class="small">The {len(a["records"])} records below extend the capability inventory. Existing core timing is unchanged. Setup is based on official guidance; checks and examples are proposed, not executed.</p>'+''.join(record_html(r,sources,prefix) for r in a['records'])+source_html(list(dict.fromkeys(k for r in a['records'] for k in r['refs'])),sources,prefix)+'</section>'
        text0=text0.replace('<section class="section" id="evidence">',inline+'<section class="section" id="evidence">',1)
        # Accommodate either evidence tag shape; never silently omit the extension.
        if 'id="mechanisms"' not in text0:text0=text0.replace('</main>',inline+'</main>',1)
        text0=text0.replace('<a href="#evidence">Evidence</a>','<a href="#mechanisms">Mechanisms</a><a href="#evidence">Evidence</a>',1)
        text0=text0.replace('</style>', '\n'+CSS.split('@media print')[0].replace(':root{','.__unused_vars{')+'\n</style>',1) if False else text0
        # Scoped extension styling avoids changing the approved core report or slides.
        scoped='\n#mechanisms .record{border:1px solid #d9e3ec;border-radius:9px;margin:16px 0;background:#fff}#mechanisms summary{padding:20px;cursor:pointer}#mechanisms summary h3{margin:8px 0}#mechanisms .body{padding:0 22px 22px;border-top:1px solid #d9e3ec}#mechanisms .label{font-size:12px;text-transform:uppercase;letter-spacing:.06em;font-weight:700;margin:20px 0 8px}#mechanisms .split{display:grid;grid-template-columns:1fr 1fr;gap:22px}#mechanisms .box{padding:16px;background:#f3f6f9}#mechanisms .source{padding:16px;margin:12px 0;background:#fff;border:1px solid #d9e3ec;overflow-wrap:anywhere}#mechanisms .refs a{display:inline-block;margin:8px 14px 0 0}#mechanisms .badge{font-size:12px;padding:4px 8px;background:#eef3f8;border-radius:4px;display:inline-block;margin:4px}#mechanisms .limit{border-left:3px solid #ad7415;padding-left:15px}@media(max-width:760px){#mechanisms .split{grid-template-columns:1fr}}'
        text0=text0.replace('</style>',scoped+'</style>',1);path.write_text(text0)
        md=root/key/'report.md';md.write_text(md.read_text()+f'\n\n## Mechanisms and enablement extension\n\nSee the [detailed source-linked setup and verification guide](../mechanisms/{key}.md). It contains {len(a["records"])} records, a proposed six-step exercise and lifecycle controls. Existing presentation timing is unchanged.\n')
    put('mechanisms/index.html',overview_html(data));put('mechanisms/overview.html',overview_html(data))
    omd=['# Platform mechanisms and enablement',data_boundary(),'## Comparison']
    for row in data['comparison']:omd += ['### '+row[0]]+[f'**{lab}:** {val}' for lab,val in zip(['Context','Activation','Surface','Verification','Limit'],row[1:])]
    omd+=['## Proposed common controls']
    for title,body,ids in data['controls']:omd += ['### '+title,body]+[f'[{k}]({sources[k]["url"]})' for k in ids]
    omd+=['## Evidence qualifications']+data['evidence_notes'];put('mechanisms/overview.md','\n\n'.join(omd)+'\n')
    counts['overview'],trans=overview_ppt(data,root/'presentations/overview-mechanisms-reference.pptx');put('mechanisms/overview-presentation-notes.md',trans)
    # Architecture gets the comparison explicitly, not a fictitious shared implementation.
    path=root/'architecture/common-patterns.html';txt=path.read_text();sect='<section class="section" id="mechanisms"><div class="kicker">Mechanism comparison / documented product boundaries</div><h2>Do not standardize away the product-specific controls.</h2><p>The common review sequence does not imply a shared memory store, execution runtime or permission model.</p><div class="actions"><a class="button" href="../mechanisms/index.html">Mechanism comparison and setup guides</a><a class="button" href="../presentations/overview-mechanisms-reference.pptx">Reference PowerPoint</a></div>'
    for row in data['comparison']:sect+='<article class="panel" style="margin:20px 0"><h3>'+H(row[0])+'</h3>'+p(row[1])+p('Activation: '+row[2])+p('Verification: '+row[4])+p('Boundary: '+row[5])+'</article>'
    sect+='</section>';txt=txt.replace('</main>',sect+'</main>',1).replace('<a href="#evidence">Evidence</a>','<a href="#mechanisms">Mechanisms</a><a href="#evidence">Evidence</a>',1);path.write_text(txt)
    p0=root/'architecture/report.md';p0.write_text(p0.read_text()+'\n\n## Mechanism comparison\n\nSee [product-specific storage, configuration and verification mechanisms](../mechanisms/overview.md). No shared runtime or memory store is implied.\n')
    index=root/'index.html';it=index.read_text();entry='<section class="section" id="mechanism-guides"><div class="kicker">Mechanisms and enablement</div><h2>Configure and verify the documented behavior.</h2><p>32 mechanism records across four platforms, with official setup steps, correction/stop controls and proposed independent checks. Five additional reference decks are untimed. Four optional exercises add 44 minutes; the original core remains 160 minutes.</p><a class="button" href="mechanisms/index.html">Open the mechanism guide collection</a></section>'
    it=it.replace('</main>',entry+'</main>',1);index.write_text(it)
    p0=root/'README.md';p0.write_text(p0.read_text()+'\n## Mechanisms and enablement\n\n[Open the mechanism guide index](mechanisms/index.html), or read the [comparison](mechanisms/overview.md). Four platform guides contain 32 documented mechanism records, official enablement procedures, proposed checks and lifecycle controls. Five reference PowerPoints are untimed. Optional prepared-result exercises total 44 minutes, separate from the unchanged 160-minute core programme.\n\nBuild source: `content/mechanisms.json` and `tools/build_mechanisms.py`. See [mechanism validation](mechanisms/VALIDATION.md) for source qualifications and test boundaries.\n')
    put('mechanisms/VALIDATION.md','# Mechanism review boundary\n\nSource register: 46 official AWS/Kiro pages. Original 44-source mechanism review dated 21 October 2026; this follow-up adds two references and rechecks selected critical mechanisms, not every historical claim. These guides document visible behavior, configuration and data lifecycle; they do not infer undisclosed internals or claim customer-controlled model training.\n\n## Source qualifications\n\n'+'\n\n'.join(data['evidence_notes'])+'\n\n## Verification\n\nAll independent checks, example prompts and six-step exercises are proposed. No live AWS account, agent run, security test or production change was executed. Structural source/reference/timing checks are in `../qa/mechanism-validation.json`. Browser and rendered-slide checks are separate; their absence is not a pass. The original core programme remains 160 minutes including 48 demo minutes. Mechanism reference slides are untimed; optional exercises add 44 minutes.\n')
    put('mechanisms/source-register.json',json.dumps(sources,ensure_ascii=False,indent=2)+'\n')
    result={'revision':data['revision'],'records':sum(len(a['records']) for a in data['agents'].values()),'official_sources':len(sources),'core_minutes_unchanged':160,'core_demo_minutes_unchanged':48,'optional_exercise_minutes':sum(a['use_case']['minutes'] for a in data['agents'].values()),'reference_pptx_slides':counts,'errors':[],'boundary':'Structural checks, not live product tests or measured outcomes.'}
    for key,a in data['agents'].items():
        assert sum(s[-1] for s in a['use_case']['steps'])==a['use_case']['minutes']*60
        for r in a['records']:assert all(x in sources for x in r['refs'])
    put('qa/mechanism-validation.json',json.dumps(result,indent=2)+'\n')
    put('PUBLISHING.md','# Publishing and rebuilding\n\nThe editorial-format revision was published through the successful [approved editorial import](https://github.com/lyrel-dev/frontier-agents/actions/runs/39575904010). The mechanism extension is generated from `content/mechanisms.json`; verify its publication by the latest successful documentation workflow and the presence of `mechanisms/index.html` on `main`.\n\nRepository visibility is preserved. Repository storage is not a GitHub Pages deployment. The source and generated files are committed one changed document per commit; unrelated history is retained.\n\n## Rebuild\n\n```sh\npython tools/assemble_content.py\npython tools/build_suite.py\npython tools/build_suite.py --check\npython tools/editorial_lint.py\n```\n\nThe existing documentation workflow runs when `documentation-build.json` is updated. Workflow permissions and definitions are not changed by the mechanism extension. The build checks references, local links and core timing; browser and rendered-slide checks are recorded separately. No live AWS execution or measured benefit is claimed.\n')

    for rel,extra in [('CHANGELOG.md','\n\n## Mechanisms and enablement extension\n\nAdded 32 mechanism records, four platform setup guides, a comparison guide, five untimed reference decks, six-step optional exercises and a dated source register. Retained all core presentation budgets.\n'),('VALIDATION.md','\n\n## Mechanisms extension\n\nSee [the mechanism-specific validation boundary](mechanisms/VALIDATION.md) and `qa/mechanism-validation.json`. This extension uses a new official-documentation review, with explicit source discrepancies, and does not claim a new audit of every earlier report statement.\n')]:
        p0=root/rel;p0.write_text(p0.read_text()+extra)
    print(json.dumps(result,indent=2))
if __name__=='__main__':build_mechanisms(Path(__file__).resolve().parents[1])
