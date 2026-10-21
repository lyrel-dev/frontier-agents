"""Native, editable PowerPoint publication edition of the report slide source."""
from __future__ import annotations
import math, textwrap
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

INK='182B43'; MUTED='55657A'; NAVY='10243B'; LINE='DBE2E9'; SOFT='F1F5F8'; WHITE='FFFFFF'
FONT='Arial'

def box(slide,x,y,w,h,fill=WHITE,line=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill)
    if line:shape.line.color.rgb=RGBColor.from_string(line);shape.line.width=Pt(.8)
    else:shape.line.fill.background()
    if radius:
        try:shape.adjustments[0]=.08
        except (IndexError,AttributeError):pass
    return shape

def fitted_size(value,w,h,size,minimum=13):
    for pt in [size-i*.5 for i in range(max(1,int((size-minimum)*2)+1))]:
        width=max(8,int((w-.03)*72/(pt*.51)))
        lines=sum(max(1,len(textwrap.wrap(p,width=width,break_long_words=True))) for p in str(value).split('\n'))
        if lines*pt*1.17/72 <= h-.015:return pt
    return minimum

def text(slide,value,x,y,w,h,size=18,color=INK,bold=False,minimum=13,link=None):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame
    tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    pt=fitted_size(value,w,h,size,minimum)
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(pt);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(3);p.line_spacing=1.12
        if link:
            for run in p.runs:run.hyperlink.address=link
    return shape

def stamp(seconds):return f'{int(seconds)//60:02}:{int(seconds)%60:02}'

def base(prs,deck,slide_data,index,start,accent):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
    box(slide,0,0,13.333,.12,accent)
    text(slide,f'{deck["title"].upper()}  /  {slide_data["phase"].upper()}',.55,.34,11.9,.25,10.5,accent,True,minimum=10.5)
    text(slide,slide_data['title'],.55,.85,12.05,.88,31,INK,True,minimum=25)
    text(slide,slide_data['takeaway'],.55,1.78,12.05,.55,18,MUTED,minimum=15)
    box(slide,.55,6.99,12.23,.01,LINE)
    text(slide,('UNTIMED APPENDIX' if slide_data['phase']=='Untimed appendix' else f'{index:02} / {len(deck["slides"]):02}')+'    '+deck['title'],.55,7.12,6.3,.22,10.5,MUTED,minimum=10.5)
    text(slide,('Outside the timed session' if slide_data['phase']=='Untimed appendix' else f'{stamp(slide_data["seconds"])} allocated  |  {stamp(start)}–{stamp(start+slide_data["seconds"])}'),8.3,7.12,4.5,.22,10.5,MUTED,minimum=10.5)
    source_lines=[]
    for ref in slide_data.get('refs',[]):
        s=deck['sources'].get(ref)
        if s:source_lines.append(f'{ref} — {s["title"]}: {s["url"]}')
    notes=f'Timing: {stamp(start)}–{stamp(start+slide_data["seconds"])}; allocated {stamp(slide_data["seconds"])}.\n\n{slide_data["notes"]}\n\nOfficial evidence:\n'+'\n'.join(source_lines)
    if slide_data['kind']=='demo':
        notes+='\n\nPreflight: '+slide_data['preflight']+'\n\nPrompt: '+slide_data['prompt']+'\n\nRunbook:\n'+'\n'.join(f'{stamp(t)} — {a}: {b}' for t,a,b in slide_data['steps'])+'\n\nSuccess: '+slide_data['success']+'\n\nFallback: '+slide_data['fallback']
    slide.notes_slide.notes_text_frame.text=notes
    return slide

def card(slide,x,y,w,h,title,body,tag,accent):
    box(slide,x,y,w,h,SOFT,LINE,True)
    if tag:text(slide,tag,x+.18,y+.16,w-.36,.22,10.5,accent,True,minimum=10)
    top=y+(.52 if tag else .2)
    text(slide,title,x+.18,top,w-.36,.55,20,INK,True,minimum=16)
    text(slide,body,x+.18,top+.66,w-.36,h-(top-y)-.82,18 if h>2.4 else 16,INK,minimum=14)

def build_deck(deck:dict,story:dict,path:Path,evidence:dict,records:dict):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    prs.core_properties.title=deck['title']+' — engineering field guide'
    prs.core_properties.subject='Source-aligned publication edition; timed demonstrations and official documentation.'
    prs.core_properties.author='Frontier agents documentation'
    prs.core_properties.comments='Independent synthesis. No live AWS execution or measured benefit is claimed.'
    accent=story['accent'].lstrip('#').upper();elapsed=0
    for index,s in enumerate(deck['slides'],1):
        slide=base(prs,deck,s,index,elapsed,accent);kind=s['kind']
        if s.get('refs'):
            text(slide,'Official source IDs: '+', '.join(s['refs']),.57,6.79,11.9,.17,9.5,MUTED,minimum=9)
        if kind=='hero':
            box(slide,.55,2.55,12.23,1.05,NAVY,None,True)
            text(slide,s['outcome'],.8,2.84,11.7,.45,25,WHITE,True,minimum=21)
            for j,(a,b) in enumerate(s['metrics']):
                x=.55+j*4.14;box(slide,x,4.05,3.95,1.55,SOFT,LINE,True);text(slide,a,x+.22,4.28,3.5,.58,35,accent,True,minimum=30);text(slide,b,x+.22,4.98,3.5,.37,17,INK,minimum=14)
            text(slide,story['availability'],.57,6.02,12.1,.58,17,MUTED,minimum=14)
        elif kind in ('cards','flow','close'):
            n=len(s['items']);cols=2 if n in (2,4) else 3;rows=math.ceil(n/cols);gap=.22;w=(12.23-gap*(cols-1))/cols;h=(4.24-gap*(rows-1))/rows
            for j,item in enumerate(s['items']):card(slide,.55+(j%cols)*(w+gap),2.53+(j//cols)*(h+gap),w,h,item['title'],item['body'],item.get('tag',''),accent)
        elif kind=='lens':
            for j,(label,field) in enumerate([('Inputs','inputs'),('Agent work','actions'),('Outputs','outputs')]):card(slide,.55+j*4.14,2.53,3.95,4.24,label,'\n\n'.join(s[field]),f'{j+1:02}',accent)
        elif kind=='compare':
            for j,k in enumerate(['left','right']):card(slide,.55+j*6.23,2.53,6,3.25,s[k]['title'],s[k]['body'],s[k].get('tag',''),accent)
            box(slide,.55,5.99,12.23,.78,NAVY,None,True);text(slide,s.get('callout',''),.77,6.16,11.78,.44,17,WHITE,True,minimum=14)
        elif kind=='matrix':
            rows=s['rows'];cols=len(s['headers']);shape=slide.shapes.add_table(len(rows)+1,cols,Inches(.55),Inches(2.53),Inches(12.23),Inches(4.22));t=shape.table
            for r,row in enumerate([s['headers'],*rows]):
                for c,value in enumerate(row):
                    cell=t.cell(r,c);cell.text=value;cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(NAVY if r==0 else (WHITE if r%2==0 else SOFT));cell.margin_left=Inches(.14);cell.margin_right=Inches(.14);cell.margin_top=Inches(.09);cell.margin_bottom=Inches(.06)
                    font=fitted_size(value,12.23/cols-.28,4.22/(len(rows)+1)-.15,16,13)
                    for p in cell.text_frame.paragraphs:p.font.name=FONT;p.font.size=Pt(font);p.font.bold=r==0;p.font.color.rgb=RGBColor.from_string(WHITE if r==0 else INK);p.space_after=Pt(0)
        elif kind=='contract':
            box(slide,.55,2.53,12.23,4.24,NAVY,None,True)
            for j,(a,b) in enumerate(s['fields']):
                y=2.8+j*.61;text(slide,a,.82,y,2.8,.4,16,'AFCBE8',True,minimum=14);text(slide,b,3.7,y,8.75,.45,18,WHITE,minimum=14)
        elif kind=='demo':
            box(slide,.55,2.51,12.23,1.06,NAVY,None,True);text(slide,s['prompt'],.77,2.66,11.8,.78,15,WHITE,minimum=13)
            yy=3.75
            for label,field in [('PREFLIGHT','preflight'),('SUCCESS','success'),('FALLBACK','fallback')]:
                text(slide,label,.55,yy,3.5,.22,10.5,accent,True,minimum=10.5);text(slide,s[field],.55,yy+.28,3.65,.67,14,INK,minimum=13);yy+=1.05
            steps=s['steps'];rowh=3.04/max(1,len(steps))
            for j,(seconds,title,body) in enumerate(steps):
                y=3.73+j*rowh;box(slide,4.48,y,8.3,rowh-.06,SOFT,None,True);text(slide,stamp(seconds),4.63,y+.12,.65,.28,13,accent,True,minimum=12);text(slide,title,5.4,y+.08,7.15,.24,15,INK,True,minimum=13);text(slide,body,5.4,y+.30,7.15,rowh-.32,13.5,INK,minimum=12)
        else:raise ValueError('Unsupported slide type '+kind)
        elapsed+=s['seconds']
    assert elapsed==deck['minutes']*60
    items=evidence.get('capabilities',evidence.get('patterns',[]))
    official={x['id']:x for x in evidence['sources']}
    for ref_index,c in enumerate(items,1):
        claim=c.get('claim',c.get('rule',''))
        limit=c.get('limit',c.get('boundary',''))
        output=records[c['id']]['record']
        name=f'{c["id"]} · {c["title"]}'
        meta=' · '.join([c.get('availability',c.get('kind','Team recommendation')),c.get('delivery','Architecture pattern')])
        refslide={'title':name,'takeaway':meta,'phase':'Untimed appendix','kind':'reference','seconds':0,
                  'notes':'Reference content only; not part of the timed presentation. Product function and prerequisites come from the cited official documentation. The retention record and independent acceptance procedure are recommended team controls, not asserted native outputs. No live workflow was executed.',
                  'refs':[]}
        slide=base(prs,deck,refslide,len(deck['slides'])+ref_index,elapsed,accent)
        head='PROPOSED ARCHITECTURE PATTERN' if 'patterns' in evidence else 'DOCUMENTED FUNCTION'
        top=[(head,claim),('INVOCATION & SCOPE',c.get('setup','')),('RETAINED RECORD · PROPOSED',output)]
        for j,(title,body) in enumerate(top):
            x=.55+j*4.14;box(slide,x,2.50,3.95,2.03,SOFT,LINE,True)
            text(slide,title,x+.17,2.66,3.60,.28,10.5,accent,True,minimum=10)
            text(slide,body,x+.17,3.04,3.58,1.44,17,INK,minimum=12.5)
        box(slide,.55,4.72,6.00,1.94,SOFT,LINE,True)
        text(slide,'INDEPENDENT ACCEPTANCE · PROPOSED',.72,4.88,5.65,.27,10.5,accent,True,minimum=10)
        text(slide,c.get('check',''),.72,5.29,5.63,1.18,17,INK,minimum=12.5)
        box(slide,6.77,4.72,6.01,1.94,'FCF9F1',LINE,True)
        text(slide,'KNOWN LIMITS AND EFFECTS',6.94,4.88,5.67,.27,10.5,'8B5B1C',True,minimum=10)
        text(slide,limit,6.94,5.29,5.63,1.18,17,INK,minimum=12.5)
        source_lines=[]
        for rid in [*c.get('refs',[])]:
            source=official.get(rid)
            if source:source_lines.append(f'{rid} — {source["title"]}: {source["url"]}')
        for imp in c.get('implementations',[]):
            if isinstance(imp,dict):
                for rid in imp.get('refs',[]):
                    source=official.get(rid)
                    if source:source_lines.append(f'{rid} — {source["title"]}: {source["url"]}')
        if not source_lines and isinstance(c.get('implementation'),dict):
            for imp in c['implementation'].values():
                if isinstance(imp,dict):
                    for rid in imp.get('refs',[]):
                        source=official.get(rid)
                        if source:source_lines.append(f'{rid} — {source["title"]}: {source["url"]}')
        slide.notes_slide.notes_text_frame.text += '\n\nCapability evidence:\n'+'\n'.join(dict.fromkeys(source_lines))
        source_ids=[entry.split(' — ')[0] for entry in dict.fromkeys(source_lines)]
        if source_ids:
            text(slide,'Official source IDs: '+', '.join(source_ids[:9])+(' + others in notes' if len(source_ids)>9 else ''),.57,6.79,11.9,.17,9.5,MUTED,minimum=9)
    appendix={'title':'Official evidence and delivery boundary','takeaway':'Use the linked product documentation and the complete report source register.','phase':'Untimed appendix','kind':'matrix','seconds':0,'notes':'This appendix is outside the timed session. All source URLs are listed below and in the relevant slide notes. No live workflow or measured benefit is claimed.','refs':list(deck['sources'])}
    slide=base(prs,deck,appendix,len(deck['slides'])+1,elapsed,accent)
    for j,(ref,source) in enumerate(list(deck['sources'].items())[:6]):
        y=2.55+j*.58;text(slide,f'{ref}  {source["title"]}',.65,y,11.95,.35,18,accent,True,minimum=14,link=source['url'])
    text(slide,'Complete source URLs are embedded in speaker notes. Demo slots include fallback walkthroughs; questions and breaks are additional.',.65,6.19,11.95,.52,15,MUTED,minimum=13)
    prs.save(path)
